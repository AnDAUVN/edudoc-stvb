import docx

doc = docx.Document('storage/uploads/temp_2100001234_01_2100001234_NGUYỄNVĂNA_CongVanDeNghi.docx')
cell = doc.tables[1].cell(0, 0)
for i, p in enumerate(cell.paragraphs):
    print(f"P{i}: text='{p.text}'")
    print(f"   style: {p.style.name if p.style else None}")
    pPr = getattr(p._element, 'pPr', None)
    if pPr is not None:
        rPr_p = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
        if rPr_p is not None:
            sz_p = rPr_p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
            print(f"   pPr/rPr sz: {float(sz_p.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'))/2.0 if sz_p is not None else None}")
    print(f"   runs count: {len(p.runs)}")
    for j, r in enumerate(p.runs):
        sz = r.font.size.pt if r.font.size else None
        rPr = getattr(r._element, 'rPr', None)
        sz_xml = None
        szCs_xml = None
        if rPr is not None:
            sz_elem = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
            if sz_elem is not None:
                sz_xml = float(sz_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')) / 2.0
            szCs_elem = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}szCs')
            if szCs_elem is not None:
                szCs_xml = float(szCs_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')) / 2.0
        print(f"     r{j}: text='{r.text}', sz={sz}, sz_xml={sz_xml}, szCs_xml={szCs_xml}")
