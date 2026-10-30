import os
import difflib
import hmac
import hashlib
import datetime
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor

SECRET_KEY = "STVB_NTTU_EDUCATIONAL_SECRET_2026"

ZERO_BIN_0 = '\u200b'  # Zero-width space -> bit 0
ZERO_BIN_1 = '\u200c'  # Zero-width non-joiner -> bit 1

def generate_anti_cheat_token(sbd, mssv, doc_type, timestamp):
    """Tạo chữ ký số bảo mật cho sinh viên"""
    message = f"{sbd}_{mssv}_{doc_type}_{timestamp}"
    token = hmac.new(SECRET_KEY.encode('utf-8'), message.encode('utf-8'), hashlib.sha256).hexdigest()
    return token

def encode_hidden_mssv_to_text(text, mssv):
    """
    Cấy mã MSSV của sinh viên dưới dạng ký tự vô hình vào giữa nội dung văn bản
    """
    binary = ''.join(format(ord(c), '08b') for c in str(mssv))
    zero_width_str = binary.replace('0', ZERO_BIN_0).replace('1', ZERO_BIN_1)
    
    # Cấy vào sau từ đầu tiên
    words = text.split(' ', 1)
    if len(words) >= 2:
        return words[0] + zero_width_str + ' ' + words[1]
    return text + zero_width_str

def plant_watermark_in_line(line, mssv):
    """
    Cấy thủy ấn số MSSV vào dòng văn bản ở các vị trí chiến lược
    đảm bảo người dùng định dạng / cắt dán / xóa nhãn vẫn giữ nguyên thủy ấn.
    """
    binary = ''.join(format(ord(c), '08b') for c in str(mssv))
    zero_width_str = binary.replace('0', ZERO_BIN_0).replace('1', ZERO_BIN_1)
    
    res = line
    # Nếu là dòng hướng dẫn trong ngoặc vuông, cấy vào cuối dòng để không làm vỡ cụm từ kiểm thử
    if res.strip().startswith('['):
        return res + zero_width_str

    # 1. Cấy trực tiếp vào các từ khóa bất biến trong thể thức Nghị định 30
    anchors = [
        "CỘNG", "Độc", "Nơi", "Kính", "TRƯỜNG", "PHÒNG", 
        "ĐOÀN", "THÔNG", "TỜ", "ỦY BAN", "SỞ", "CÔNG TY"
    ]
    for anchor in anchors:
        if anchor in res and zero_width_str not in res:
            res = res.replace(anchor, anchor + zero_width_str, 1)
            
    # 2. Cấy thêm vào sau từ đầu tiên của dòng
    words = res.split(' ', 1)
    if len(words) >= 2 and zero_width_str not in words[0]:
        res = words[0] + zero_width_str + ' ' + words[1]
    elif zero_width_str not in res:
        res = res + zero_width_str
        
    return res

def extract_all_hidden_mssvs(text):
    """
    Trích xuất và giải mã toàn bộ các thủy ấn số vô hình có trong văn bản
    """
    pattern = f"[{ZERO_BIN_0}{ZERO_BIN_1}]{{16,}}"
    matches = re.findall(pattern, text)
    if not matches:
        return []
    
    found_mssvs = []
    for candidate in matches:
        raw_binary = candidate.replace(ZERO_BIN_0, '0').replace(ZERO_BIN_1, '1')
        chars = []
        for i in range(0, len(raw_binary), 8):
            byte = raw_binary[i:i+8]
            if len(byte) == 8:
                try:
                    val = int(byte, 2)
                    if 32 <= val <= 126:  # Ký tự ASCII in được
                        chars.append(chr(val))
                    else:
                        break
                except:
                    break
        decoded = ''.join(chars)
        if len(decoded) >= 4 and any(c.isdigit() for c in decoded):
            if decoded not in found_mssvs:
                found_mssvs.append(decoded)
            
    return found_mssvs

def create_student_raw_template(sbd, mssv, student_name, doc_type_id, raw_text, output_path):
    """
    Khởi tạo file docx thô:
    - Chưa căn lề Nghị định 30 (để lề mặc định 1 inch)
    - Chưa chỉnh font Times New Roman (để Calibri 11pt, đen, không đậm/nghiêng)
    - Cấy mã token bảo mật và thủy ấn số MSSV
    """
    doc = Document()
    
    # 1. Để lề mặc định 1 inch (2.54 cm) - Sinh viên bắt buộc phải tự chỉnh sang NĐ 30
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # 2. Tạo Token bảo mật & Ghi vào Metadata ẩn của file Word
    now_str = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
    token = generate_anti_cheat_token(sbd, mssv, doc_type_id, now_str)
    
    core_props = doc.core_properties
    core_props.identifier = f"STVB_SIGNATURE:{token}"
    core_props.keywords = f"{sbd}|{mssv}|{student_name}|{doc_type_id}|{now_str}"
    core_props.author = f"STVB Educational System - NTTU"
    core_props.comments = f"Bai tap STVB - Sinh vien: {student_name} - MSSV: {mssv}"
    
    # 3. Ghi nội dung thô kèm thủy ấn vô hình
    lines = raw_text.strip().split("\n")
    for idx, line in enumerate(lines):
        line_to_write = plant_watermark_in_line(line, mssv)
            
        p = doc.add_paragraph()
        run = p.add_run(line_to_write)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.bold = False
        run.font.italic = False
        run.font.color.rgb = RGBColor(0, 0, 0)
        
        # Chưa căn lề Justify, không thụt đầu dòng
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

    doc.save(output_path)
    return token


def compute_box6_similarity(text1: str, text2: str) -> float:
    """
    Tính độ tương đồng nội dung văn bản Ô 6:
    Kết hợp SequenceMatcher theo cấp độ từ (word-level) và N-gram overlap (cụm 3 từ).
    Trả về tỷ lệ % từ 0.0 đến 100.0
    """
    if not text1 or not text2:
        return 0.0

    # Làm sạch và tách từ (loại bỏ dấu câu, chuyển chữ thường)
    w1 = [w.strip(".,;:?!-()\"'“”‘’+*") for w in text1.lower().split() if len(w.strip(".,;:?!-()\"'“”‘’+*")) > 0]
    w2 = [w.strip(".,;:?!-()\"'“”‘’+*") for w in text2.lower().split() if len(w.strip(".,;:?!-()\"'“”‘’+*")) > 0]

    if not w1 or not w2:
        return 0.0

    seq = difflib.SequenceMatcher(None, w1, w2)
    seq_ratio = seq.ratio()

    if len(w1) >= 3 and len(w2) >= 3:
        ng1 = set(tuple(w1[i:i+3]) for i in range(len(w1)-2))
        ng2 = set(tuple(w2[i:i+3]) for i in range(len(w2)-2))
        ngram_overlap = len(ng1 & ng2) / max(len(ng1), len(ng2))
    else:
        ngram_overlap = seq_ratio

    final_score = max(seq_ratio, ngram_overlap)
    return round(final_score * 100, 1)

def extract_box6_from_docx(doc_or_path) -> str:
    """Trích xuất toàn bộ văn bản phần thân bài (Ô 6) ngoài bảng"""
    if isinstance(doc_or_path, str):
        if not os.path.exists(doc_or_path):
            return ""
        doc = Document(doc_or_path)
    else:
        doc = doc_or_path

    paras = []
    in_kinh_gui = False

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            in_kinh_gui = False
            continue

        t_clean = re.sub(r'[\u200b\u200c\u200d\ufeff]', '', t).strip()
        if not t_clean:
            continue

        # Bỏ qua các dòng nhãn hướng dẫn trong phôi đề
        if t_clean.startswith('[') and t_clean.endswith(']'):
            continue
        if any(t_clean.lower().startswith(kw) for kw in [
            'cơ quan chủ quản', 'cơ quan ban hành', 'quốc hiệu', 'tiêu ngữ',
            'số ký hiệu', 'địa danh ngày tháng', 'trích yếu', 'tên loại'
        ]):
            continue

        if 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA' in t_clean or 'Độc lập - Tự do - Hạnh phúc' in t_clean:
            continue
        if re.match(r'^Số\s*:', t_clean, re.IGNORECASE) or t_clean.startswith('SỐ:'):
            continue
        if 'ngày' in t_clean.lower() and 'tháng' in t_clean.lower() and 'năm' in t_clean.lower() and len(t_clean) < 60:
            continue
        if t_clean.startswith(('V/v', 'v/v', 'Về việc', 'về việc')) and len(t_clean) < 80:
            continue

        # Bỏ qua Kính gửi
        if t_clean.lower().startswith('kính gửi'):
            in_kinh_gui = True
            continue

        if in_kinh_gui:
            if t_clean.startswith(('-', '–', '—', '+')) or t_clean.endswith(';') or t_clean.endswith('.'):
                if any(t_clean.startswith(kw) for kw in ['Hiện nay', 'Căn cứ', 'Thực hiện', 'Nhằm', 'Để ']):
                    in_kinh_gui = False
                else:
                    continue
            else:
                in_kinh_gui = False

        if t_clean.startswith('Nơi nhận') or t_clean.startswith('Lưu:') or t_clean == '- Như kính gửi;' or t_clean == '- Lưu: VT.':
            continue
        if any(kw in t_clean.upper() for kw in ['HIỆU TRƯỞNG', 'GIÁM ĐỐC', 'TRƯỞNG PHÒNG', 'TRƯỞNG KHOA', 'BÍ THƯ']):
            continue
        # Bỏ qua dòng họ tên người ký nếu là chuỗi ngắn 2-4 từ không dấu câu
        if len(t_clean.split()) in [2, 3, 4] and not any(punc in t_clean for punc in ['.', ',', ';', '!', '?', ':']):
            if t_clean in ['Trần Ái Cầm', 'Nguyễn Văn An', 'Lê Văn Bình', 'Nguyễn Hoàng Long', 'Nguyễn Văn A']:
                continue

        paras.append(t_clean)

    return '\n'.join(paras)

def check_box6_plagiarism(docx_path, current_sbd, current_mssv, current_name, doc_type=None, class_code=None, students_db=None):
    """
    Kiểm tra đạo văn Thân bài (Ô 6):
    - Đối chiếu với các bài đã lưu trong SQLite box6_submissions
    - Đối chiếu với tất cả file bài làm trong storage/uploads/ và storage/submitted/
    - Đối chiếu với các mẫu văn bản khác trong document_templates.json
    """
    try:
        curr_box6 = extract_box6_from_docx(docx_path)
        words = [w for w in curr_box6.split() if w]
        if len(words) < 15:
            return {"is_fraud": False, "box6_text": curr_box6}

        from engine.common import DATA_DIR, load_json_file, STORAGE_UPLOADS, STORAGE_SUBMITTED
        from engine.db import get_all_box6_submissions, save_box6_submission

        # 1. Kiểm tra đối chiếu với các mẫu văn bản khác trong hệ thống
        templates_path = os.path.join(DATA_DIR, "document_templates.json")
        all_tmpls = load_json_file(templates_path, {})
        curr_tmpl = all_tmpls.get(doc_type, {})
        is_curr_blank = curr_tmpl.get("is_body_blank", True)

        if is_curr_blank:
            for other_id, other_tmpl in all_tmpls.items():
                if other_id == doc_type:
                    continue
                other_body = other_tmpl.get("body_part", "")
                if other_body and len(other_body.split()) >= 15:
                    sim = compute_box6_similarity(curr_box6, other_body)
                    if sim >= 70.0:
                        return {
                            "is_fraud": True,
                            "fraud_type": "COPY_NOI_DUNG_BEN_NGOAI",
                            "reason": "Sao chép nội dung từ mẫu văn bản khác",
                            "student_a": {
                                "sbd": current_sbd,
                                "mssv": current_mssv,
                                "name": current_name,
                                "class_code": class_code or "Chung"
                            },
                            "similarity": sim,
                            "evidence": f"Sinh viên {current_name} (MSSV: {current_mssv}) đã copy nội dung Thân bài từ mẫu '{other_tmpl.get('name', other_id)}' dán vào bài làm! (Độ trùng lặp: {sim}%)"
                        }

        # 2. Kiểm tra đối chiếu với các bài làm của sinh viên khác trong SQLite box6_submissions
        other_submissions = get_all_box6_submissions(doc_type=doc_type, exclude_mssv=current_mssv)
        for other in other_submissions:
            other_text = other.get("box6_text", "")
            if not other_text:
                continue
            sim = compute_box6_similarity(curr_box6, other_text)
            if sim >= 75.0:
                target_name = other.get("name") or f"Sinh viên (MSSV: {other['mssv']})"
                if students_db and str(other['mssv']) in students_db:
                    target_name = students_db[str(other['mssv'])].get("name", target_name)
                return {
                    "is_fraud": True,
                    "fraud_type": "COPY_PASTE_NOI_DUNG_O6",
                    "reason": "Sao chép nội dung Thân bài (Ô 6) từ bạn học",
                    "student_a": {
                        "sbd": current_sbd,
                        "mssv": current_mssv,
                        "name": current_name,
                        "class_code": class_code or "Chung"
                    },
                    "student_b": {
                        "mssv": other["mssv"],
                        "name": target_name,
                        "class_code": other.get("class_code", "")
                    },
                    "similarity": sim,
                    "evidence": f"Sinh viên A ({current_name} - MSSV: {current_mssv}) đã sao chép {sim}% nội dung Thân bài (Ô 6) từ bài làm của Sinh viên B ({target_name} - MSSV: {other['mssv']})."
                }

        # 3. Quét kiểm tra trực tiếp các file docx khác trong storage/uploads/ và storage/submitted/
        scan_dirs = [STORAGE_UPLOADS, STORAGE_SUBMITTED]
        for sdir in scan_dirs:
            if not os.path.exists(sdir):
                continue
            for fn in os.listdir(sdir):
                if not fn.endswith(".docx") or fn.startswith("~"):
                    continue
                # Bỏ qua chính file hiện tại của sinh viên này
                if str(current_mssv) in fn:
                    continue
                fpath = os.path.join(sdir, fn)
                try:
                    other_doc = Document(fpath)
                    props = other_doc.core_properties
                    kw = getattr(props, "keywords", "") or ""
                    kw_parts = kw.split("|")
                    if len(kw_parts) >= 3:
                        other_mssv = kw_parts[1]
                        other_name = kw_parts[2]
                    else:
                        m_mssv = re.search(r'\b(\d{10})\b', fn)
                        other_mssv = m_mssv.group(1) if m_mssv else "Bạn cùng lớp"
                        other_name = f"Sinh viên (MSSV: {other_mssv})"

                    if str(other_mssv).strip() == str(current_mssv).strip():
                        continue

                    other_text = extract_box6_from_docx(other_doc)
                    if other_text and len(other_text.split()) >= 15:
                        sim = compute_box6_similarity(curr_box6, other_text)
                        if sim >= 75.0:
                            # Lưu vào db để các truy vấn sau nhanh hơn
                            save_box6_submission(other_mssv, other_name, class_code or "Chung", doc_type or "", other_text)
                            return {
                                "is_fraud": True,
                                "fraud_type": "COPY_PASTE_NOI_DUNG_O6",
                                "reason": "Sao chép nội dung Thân bài (Ô 6) từ bạn học",
                                "student_a": {
                                    "sbd": current_sbd,
                                    "mssv": current_mssv,
                                    "name": current_name,
                                    "class_code": class_code or "Chung"
                                },
                                "student_b": {
                                    "mssv": other_mssv,
                                    "name": other_name,
                                    "class_code": class_code or "Chung"
                                },
                                "similarity": sim,
                                "evidence": f"Sinh viên A ({current_name} - MSSV: {current_mssv}) đã sao chép {sim}% nội dung Thân bài (Ô 6) từ file '{fn}' của Sinh viên B ({other_name} - MSSV: {other_mssv})."
                            }
                except Exception:
                    continue

        return {"is_fraud": False, "box6_text": curr_box6}
    except Exception as e:
        return {"is_fraud": False, "box6_text": ""}

def check_anti_cheat(docx_path, current_sbd, current_mssv, current_name, students_db=None, doc_type=None, class_code=None):
    """
    Kiểm tra gian lận:
    - Chiêu 1: SV A nộp nguyên file của SV B (Đổi tên file)
    - Chiêu 2: SV A mở file mình nhưng Copy-Paste nội dung từ file của SV B
    - Chiêu 3: SV copy toàn bộ nội dung từ tài liệu bên ngoài dán đè vào file mẫu
    """
    try:
        doc = Document(docx_path)
        core_props = doc.core_properties
        
        # KIỂM TRA 1: METADATA & TOKEN
        token_header = getattr(core_props, "identifier", "") or ""
        keywords = getattr(core_props, "keywords", "") or ""
        
        if not token_header.startswith("STVB_SIGNATURE:") or not keywords:
            return {
                "is_fraud": True,
                "fraud_type": "FILE_KHONG_HOP_LE",
                "reason": "File không có chữ ký số STVB",
                "message": "File bài làm không chứa mã bảo mật hợp lệ của hệ thống! Sinh viên phải sử dụng đúng file mẫu tải về từ phần mềm."
            }
            
        parts = keywords.split("|")
        if len(parts) >= 5:
            orig_sbd, orig_mssv, orig_name, orig_doc_type, timestamp = parts[0], parts[1], parts[2], parts[3], parts[4]
            
            # Nếu MSSV trong file không khớp với MSSV người đang đăng nhập nộp bài
            if str(orig_mssv).strip() != str(current_mssv).strip():
                return {
                    "is_fraud": True,
                    "fraud_type": "NOP_NGUYEN_FILE_BAN",
                    "reason": "Nộp nguyên file của bạn học",
                    "student_a": {
                        "sbd": current_sbd,
                        "mssv": current_mssv,
                        "name": current_name
                    },
                    "student_b": {
                        "sbd": orig_sbd,
                        "mssv": orig_mssv,
                        "name": orig_name
                    },
                    "evidence": f"Sinh viên A ({current_name} - MSSV: {current_mssv}) đã lấy nguyên file gốc được cấp phát cho Sinh viên B ({orig_name} - MSSV: {orig_mssv}) để nộp."
                }
                
        # KIỂM TRA 2: THỦY ẤN VÔ HÌNH TRONG VĂN BẢN (COPY - PASTE NỘI DUNG)
        full_text = "\n".join([p.text for p in doc.paragraphs])
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text
                    
        all_found_mssvs = extract_all_hidden_mssvs(full_text)
        
        # 2a. Phát hiện mã thủy ấn của sinh viên khác -> Chép bài từ bạn học
        stolen_mssvs = [m for m in all_found_mssvs if str(m).strip() != str(current_mssv).strip()]
        
        if stolen_mssvs:
            target_b_mssv = stolen_mssvs[0]
            # Tra cứu thông tin sinh viên B theo MSSV từ Database nếu có
            student_b_name = "Sinh viên lớp (MSSV: " + str(target_b_mssv) + ")"
            if students_db and str(target_b_mssv) in students_db:
                student_b_name = students_db[str(target_b_mssv)].get("name", student_b_name)
                
            return {
                "is_fraud": True,
                "fraud_type": "COPY_PASTE_NOI_DUNG",
                "reason": "Sao chép bài bạn học",
                "student_a": {
                    "sbd": current_sbd,
                    "mssv": current_mssv,
                    "name": current_name
                },
                "student_b": {
                    "mssv": target_b_mssv,
                    "name": student_b_name
                },
                "evidence": f"Sinh viên A ({current_name} - MSSV: {current_mssv}) đã copy-paste nội dung bài làm từ file của Sinh viên B ({student_b_name} - MSSV: {target_b_mssv}). Thủy ấn số bí mật của Sinh viên B đã bị phát hiện trong nội dung văn bản."
            }

        # 2b. Phát hiện copy nội dung bên ngoài dán đè vào file mẫu:
        # File có mã bảo mật phôi STVB nhưng toàn bộ thủy ấn MSSV của chính sinh viên đã bị xóa sạch!
        if str(current_mssv).strip() not in all_found_mssvs:
            return {
                "is_fraud": True,
                "fraud_type": "COPY_NOI_DUNG_BEN_NGOAI",
                "reason": "Sao chép toàn bộ văn bản bên ngoài",
                "student_a": {
                    "sbd": current_sbd,
                    "mssv": current_mssv,
                    "name": current_name
                },
                "evidence": f"Sinh viên {current_name} (MSSV: {current_mssv}) đã copy toàn bộ nội dung từ tài liệu bên ngoài dán đè vào file mẫu! Bài làm không còn chứa mã thủy ấn số bảo mật của phôi đề thi được hệ thống cấp phát."
            }

                # 3. KIỂM TRA ĐẠO VĂN NỘI DUNG THÂN BÀI (Ô 6) GIỮA CÁC SINH VIÊN
        effective_doc_type = doc_type or (parts[3] if 'parts' in locals() and len(parts) >= 4 else None)
        box6_result = check_box6_plagiarism(
            docx_path=docx_path,
            current_sbd=current_sbd,
            current_mssv=current_mssv,
            current_name=current_name,
            doc_type=effective_doc_type,
            class_code=class_code,
            students_db=students_db
        )
        if box6_result.get("is_fraud"):
            return box6_result

        return {"is_fraud": False, "message": "File chính chủ và bảo mật hợp lệ.", "box6_text": box6_result.get("box6_text", "")}

    except Exception as e:
        return {
            "is_fraud": True,
            "fraud_type": "ERROR_READING_FILE",
            "message": f"Không thể đọc file: {str(e)}"
        }
