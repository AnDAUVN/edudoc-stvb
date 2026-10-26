import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from engine.validator import Decree30Validator

out_path = 'storage/templates/test_cv_perfect.docx'
doc = Document()
for s in doc.sections:
    s.top_margin = Inches(2.0 / 2.54)
    s.bottom_margin = Inches(2.0 / 2.54)
    s.left_margin = Inches(3.0 / 2.54)
    s.right_margin = Inches(1.5 / 2.54)

# Table 1: 2x2
t_head = doc.add_table(rows=2, cols=2)
# Cell 0, 0: Cơ quan
p_cq = t_head.cell(0, 0).paragraphs[0]
p_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cq.paragraph_format.line_spacing = 1.0
p_cq.paragraph_format.space_after = Pt(0)
r1 = p_cq.add_run('BỘ GIÁO DỤC VÀ ĐÀO TẠO\n')
r1.font.name = 'Times New Roman'
r1.font.size = Pt(12)
r2 = p_cq.add_run('TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH')
r2.font.name = 'Times New Roman'
r2.font.size = Pt(12)
r2.bold = True

# Cell 0, 1: Quốc hiệu & Tiêu ngữ
p_qh = t_head.cell(0, 1).paragraphs[0]
p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_qh.paragraph_format.line_spacing = 1.0
p_qh.paragraph_format.space_after = Pt(0)
rqh = p_qh.add_run('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM')
rqh.font.name = 'Times New Roman'
rqh.font.size = Pt(12)
rqh.bold = True

p_tn = t_head.cell(0, 1).add_paragraph()
p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_tn.paragraph_format.line_spacing = 1.0
p_tn.paragraph_format.space_before = Pt(0)
rtn = p_tn.add_run('Độc lập - Tự do - Hạnh phúc')
rtn.font.name = 'Times New Roman'
rtn.font.size = Pt(13)
rtn.bold = True

# Cell 1, 0: Số ký hiệu & Ô 5b
p_so = t_head.cell(1, 0).paragraphs[0]
p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_so.paragraph_format.space_after = Pt(6)
rso = p_so.add_run('Số: 12/ĐHNTT-ĐT')
rso.font.name = 'Times New Roman'
rso.font.size = Pt(13)

p_5b = t_head.cell(1, 0).add_paragraph()
p_5b.alignment = WD_ALIGN_PARAGRAPH.CENTER
r5b = p_5b.add_run('V/v đề nghị hỗ trợ kinh phí nghiên cứu')
r5b.font.name = 'Times New Roman'
r5b.font.size = Pt(12)

# Cell 1, 1: Địa danh ngày tháng
p_dd = t_head.cell(1, 1).paragraphs[0]
p_dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
rdd = p_dd.add_run('Thành phố Hồ Chí Minh, ngày 15 tháng 01 năm 2026')
rdd.font.name = 'Times New Roman'
rdd.font.size = Pt(13)
rdd.italic = True

# Ô 9a: Kính gửi
p_kg = doc.add_paragraph()
p_kg.paragraph_format.space_before = Pt(12)
p_kg.paragraph_format.space_after = Pt(6)
rkg = p_kg.add_run('Kính gửi: Bộ Khoa học và Công nghệ.')
rkg.font.name = 'Times New Roman'
rkg.font.size = Pt(13)

# Ô 6: Thân bài
p_body = doc.add_paragraph()
p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_body.paragraph_format.first_line_indent = Inches(1.27 / 2.54)
p_body.paragraph_format.space_after = Pt(6)
p_body.paragraph_format.line_spacing = 1.15
rbody = p_body.add_run('Thực hiện Kế hoạch nghiên cứu khoa học và phát triển công nghệ năm 2026 đã được Ban Giám hiệu phê duyệt, Nhà trường kính đề nghị Bộ Khoa học và Công nghệ xem xét hỗ trợ kinh phí để triển khai dự án trọng điểm.')
rbody.font.name = 'Times New Roman'
rbody.font.size = Pt(13)

# Table 2: 1x2 Footer
t_foot = doc.add_table(rows=1, cols=2)
# Cell 0, 0: Ô 9b Nơi nhận
p_nn = t_foot.cell(0, 0).paragraphs[0]
rnn = p_nn.add_run('Nơi nhận:')
rnn.font.name = 'Times New Roman'
rnn.font.size = Pt(12)
rnn.bold = True
rnn.italic = True

p_i1 = t_foot.cell(0, 0).add_paragraph()
ri1 = p_i1.add_run('- Như kính gửi;')
ri1.font.name = 'Times New Roman'
ri1.font.size = Pt(11)

p_i2 = t_foot.cell(0, 0).add_paragraph()
ri2 = p_i2.add_run('- Ban Giám hiệu (để báo cáo);')
ri2.font.name = 'Times New Roman'
ri2.font.size = Pt(11)

p_i3 = t_foot.cell(0, 0).add_paragraph()
ri3 = p_i3.add_run('- Lưu: VT, ĐT.')
ri3.font.name = 'Times New Roman'
ri3.font.size = Pt(11)

# Cell 0, 1: Ô 7a & 7b
p_cv = t_foot.cell(0, 1).paragraphs[0]
p_cv.alignment = WD_ALIGN_PARAGRAPH.CENTER
rcv = p_cv.add_run('HIỆU TRƯỞNG')
rcv.font.name = 'Times New Roman'
rcv.font.size = Pt(13)
rcv.bold = True

for _ in range(3):
    t_foot.cell(0, 1).add_paragraph()

p_name = t_foot.cell(0, 1).add_paragraph()
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
rname = p_name.add_run('Trần Ái Cầm')
rname.font.name = 'Times New Roman'
rname.font.size = Pt(13)
rname.bold = True

doc.save(out_path)

val = Decree30Validator(out_path, 'cong_van_de_nghi')
res = val.validate_document()

print('=== KIỂM THỬ CÔNG VĂN ĐỀ NGHỊ (SỐ TRƯỚC) ===')
print(f'Passed: {res["passed_count"]} / {res["total_count"]} ({res["percentage"]}%)')
print('Can submit:', res['can_submit'])
for item in res['checklist']:
    status = '✔' if item['passed'] else '✖'
    print(f'  [{status}] {item["name"]}')
    if not item['passed']:
        for e in item['errors']:
            print(f'       - {e["issue"]}')
assert res['can_submit'] == True, 'Công văn chuẩn phải đạt 100%'
print('🎉 CÔNG VĂN ĐỀ NGHỊ VƯỢT QUA TEST 100%!')
