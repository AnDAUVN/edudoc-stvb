import requests
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. Sinh viên tạo và tải file mẫu Công văn đề nghị từ hệ thống
gen_res = requests.post("http://127.0.0.1:5000/api/template/generate", json={
    "class_code": "KN02",
    "sbd": "09",
    "mssv": "2100007777",
    "name": "Lê Hoàng Nam",
    "doc_type": "cong_van_de_nghi"
}).json()

filename = gen_res["filename"]
print("1. Generated student template (Công văn):", filename)

# 2. Sinh viên mở file này trên Word và căn chỉnh chuẩn 100% Nghị định 30
file_path = f"storage/templates/{filename}"
doc = Document(file_path)

# Xóa các đoạn text thô cũ trong body để đưa vào 2 bảng và nội dung chuẩn
for p in list(doc.paragraphs):
    p._element.getparent().remove(p._element)

# Căn lề A4 chuẩn NĐ 30: Trên 2cm, Dưới 2cm, Trái 3cm, Phải 1.5cm
for s in doc.sections:
    s.top_margin = Inches(2.0 / 2.54)
    s.bottom_margin = Inches(2.0 / 2.54)
    s.left_margin = Inches(3.0 / 2.54)
    s.right_margin = Inches(1.5 / 2.54)

# Bảng 1: Header 2x2
t_head = doc.add_table(rows=2, cols=2)
# Cell 0, 0: Cơ quan
p_cq = t_head.cell(0, 0).paragraphs[0]
p_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cq.paragraph_format.line_spacing = 1.0
p_cq.paragraph_format.space_after = Pt(0)
r1 = p_cq.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
r1.font.name = "Times New Roman"
r1.font.size = Pt(12)
r2 = p_cq.add_run("TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH")
r2.font.name = "Times New Roman"
r2.font.size = Pt(12)
r2.bold = True

# Cell 0, 1: Quốc hiệu & Tiêu ngữ
p_qh = t_head.cell(0, 1).paragraphs[0]
p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_qh.paragraph_format.line_spacing = 1.0
p_qh.paragraph_format.space_after = Pt(0)
rqh = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
rqh.font.name = "Times New Roman"
rqh.font.size = Pt(12)
rqh.bold = True

p_tn = t_head.cell(0, 1).add_paragraph()
p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_tn.paragraph_format.line_spacing = 1.0
p_tn.paragraph_format.space_before = Pt(0)
rtn = p_tn.add_run("Độc lập - Tự do - Hạnh phúc")
rtn.font.name = "Times New Roman"
rtn.font.size = Pt(13)
rtn.bold = True

# Cell 1, 0: Số ký hiệu & Ô 5b (Trích yếu V/v của Công văn nằm trong bảng)
p_so = t_head.cell(1, 0).paragraphs[0]
p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_so.paragraph_format.space_after = Pt(6)
rso = p_so.add_run("Số: 88/ĐHNTT-ĐT")
rso.font.name = "Times New Roman"
rso.font.size = Pt(13)

p_5b = t_head.cell(1, 0).add_paragraph()
p_5b.alignment = WD_ALIGN_PARAGRAPH.CENTER
r5b = p_5b.add_run("V/v đề nghị hỗ trợ kinh phí nghiên cứu khoa học")
r5b.font.name = "Times New Roman"
r5b.font.size = Pt(12)

# Cell 1, 1: Địa danh ngày tháng
p_dd = t_head.cell(1, 1).paragraphs[0]
p_dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
rdd = p_dd.add_run("Thành phố Hồ Chí Minh, ngày 15 tháng 01 năm 2026")
rdd.font.name = "Times New Roman"
rdd.font.size = Pt(13)
rdd.italic = True

# Ô 9a: Kính gửi (Công văn có Kính gửi đặt trước thân bài)
p_kg = doc.add_paragraph()
p_kg.paragraph_format.space_before = Pt(12)
p_kg.paragraph_format.space_after = Pt(6)
rkg = p_kg.add_run("Kính gửi: Bộ Khoa học và Công nghệ.")
rkg.font.name = "Times New Roman"
rkg.font.size = Pt(13)

# Ô 6: Thân bài (Justify, thụt 1.27cm, giãn đoạn 6pt)
p_body = doc.add_paragraph()
p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_body.paragraph_format.first_line_indent = Inches(1.27 / 2.54)
p_body.paragraph_format.space_after = Pt(6)
p_body.paragraph_format.line_spacing = 1.15
rbody = p_body.add_run("Thực hiện Kế hoạch nghiên cứu khoa học và phát triển công nghệ năm 2026 đã được Ban Giám hiệu phê duyệt, Nhà trường kính đề nghị Bộ Khoa học và Công nghệ xem xét hỗ trợ kinh phí để triển khai dự án trọng điểm.")
rbody.font.name = "Times New Roman"
rbody.font.size = Pt(13)

# Bảng 2: Footer 1x2
t_foot = doc.add_table(rows=1, cols=2)
# Cell 0, 0: Nơi nhận Ô 9b
p_nn = t_foot.cell(0, 0).paragraphs[0]
rnn = p_nn.add_run("Nơi nhận:")
rnn.font.name = "Times New Roman"
rnn.font.size = Pt(12)
rnn.bold = True
rnn.italic = True

p_i1 = t_foot.cell(0, 0).add_paragraph()
p_i1.paragraph_format.space_after = Pt(0)
ri1 = p_i1.add_run("- Như kính gửi;")
ri1.font.name = "Times New Roman"
ri1.font.size = Pt(11)

p_i2 = t_foot.cell(0, 0).add_paragraph()
p_i2.paragraph_format.space_after = Pt(0)
ri2 = p_i2.add_run("- Ban Giám hiệu (để báo cáo);")
ri2.font.name = "Times New Roman"
ri2.font.size = Pt(11)

p_i3 = t_foot.cell(0, 0).add_paragraph()
p_i3.paragraph_format.space_after = Pt(0)
ri3 = p_i3.add_run("- Lưu: VT, ĐT.")
ri3.font.name = "Times New Roman"
ri3.font.size = Pt(11)

# Cell 0, 1: Ký tên (Hiệu trưởng ký trực tiếp)
p_cv = t_foot.cell(0, 1).paragraphs[0]
p_cv.alignment = WD_ALIGN_PARAGRAPH.CENTER
rcv = p_cv.add_run("HIỆU TRƯỞNG")
rcv.font.name = "Times New Roman"
rcv.font.size = Pt(13)
rcv.bold = True

for _ in range(3):
    t_foot.cell(0, 1).add_paragraph()

p_name = t_foot.cell(0, 1).add_paragraph()
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
rname = p_name.add_run("Trần Ái Cầm")
rname.font.name = "Times New Roman"
rname.font.size = Pt(13)
rname.bold = True

doc.save(file_path)
print("2. Student edited and formatted Công văn successfully!")

# 3. Thẩm định qua API
with open(file_path, "rb") as f:
    files = {"file": (filename, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {
        "class_code": "KN02",
        "sbd": "09",
        "mssv": "2100007777",
        "name": "Lê Hoàng Nam",
        "doc_type": "cong_van_de_nghi"
    }
    r = requests.post("http://127.0.0.1:5000/api/validate", files=files, data=data)
    res = r.json()
    print("3. Validation Status:", res.get("status"))
    print(f"   Score: {res.get('percentage')}% ({res.get('passed_count')}/{res.get('total_count')})")
    print("   Can submit:", res.get("can_submit"))
    assert res.get("can_submit") == True, "Công văn phải đạt chuẩn 100%!"

# 4. Xác nhận nộp bài
sub_res = requests.post("http://127.0.0.1:5000/api/submit", json={
    "class_code": "KN02",
    "sbd": "09",
    "mssv": "2100007777",
    "name": "Lê Hoàng Nam",
    "doc_type": "cong_van_de_nghi",
    "temp_file": res["temp_file"]
}).json()

print("4. Submission response:", sub_res.get("success"), "| Filename:", sub_res.get("submission", {}).get("filename"))
assert sub_res.get("success") == True, "Nộp bài Công văn phải thành công!"

print("\n🎉 CÔNG VĂN ĐỀ NGHỊ ĐẠT CHUẨN 100% TRÊN MÁY CHỦ LIVE!")
