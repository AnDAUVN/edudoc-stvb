import docx
import os
import glob
import shutil
import xml.etree.ElementTree as ET

target_path = 'storage/uploads/temp_2100001234_01_2100001234_NGUYỄNVĂNA_CongVanDeNghi.docx'
if not os.path.exists(target_path):
    candidates = glob.glob('storage/templates/*2100001234*CongVanDeNghi*.docx')
    if candidates:
        os.makedirs('storage/uploads', exist_ok=True)
        shutil.copy2(candidates[0], target_path)

doc = docx.Document(target_path)

def get_cascaded_font_size(paragraph, doc):
    # 1. Check runs
    for r in paragraph.runs:
        if r.font and r.font.size:
            return round(r.font.size.pt, 1)
        rPr = getattr(r._element, 'rPr', None)
        if rPr is not None:
            sz = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
            if sz is not None and sz.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'):
                return float(sz.attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val']) / 2.0
                
    # 2. Check paragraph pPr/rPr/sz
    pPr = getattr(paragraph._element, 'pPr', None)
    if pPr is not None:
        rPr = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
        if rPr is not None:
            sz = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
            if sz is not None and sz.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'):
                return float(sz.attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val']) / 2.0

    # 3. Check style and style inheritance
    style = getattr(paragraph, 'style', None)
    while style is not None:
        if style.font and style.font.size:
            return round(style.font.size.pt, 1)
        rPr = getattr(style._element, 'rPr', None)
        if rPr is not None:
            sz = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
            if sz is not None and sz.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'):
                return float(sz.attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val']) / 2.0
        style = getattr(style, 'base_style', None)

    # 4. Check docDefaults in styles.element
    try:
        styles_elem = doc.styles.element
        docDefaults = styles_elem.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}docDefaults')
        if docDefaults is not None:
            rPrDefault = docDefaults.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPrDefault')
            if rPrDefault is not None:
                rPr = rPrDefault.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
                if rPr is not None:
                    sz = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz')
                    if sz is not None and sz.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'):
                        return float(sz.attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val']) / 2.0
    except Exception as e:
        print("Error checking docDefaults:", e)

    return 13.0

if len(doc.tables) > 1:
    cell = doc.tables[1].cell(0, 0)
    for i, p in enumerate(cell.paragraphs):
        sz = get_cascaded_font_size(p, doc)
        print(f"P{i}: text='{p.text}' -> Cascaded Size = {sz} pt")
elif len(doc.tables) == 1:
    cell = doc.tables[0].cell(0, 0)
    for i, p in enumerate(cell.paragraphs):
        sz = get_cascaded_font_size(p, doc)
        print(f"P{i}: text='{p.text}' -> Cascaded Size = {sz} pt")
else:
    print("Document has no tables, cascaded font size check skipped.")
