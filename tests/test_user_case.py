import os
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from engine.validator import Decree30Validator

def test_user_screenshot_case():
    doc = Document()
    
    # Tạo Table 1 hàng 2 cột giống hệt ảnh chụp màn hình của người dùng
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Cột trái: CÔNG TY TNHH HẢI THÀNH (13pt, Bold)
    cell_left = tbl.cell(0, 0)
    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_l = p_left.add_run("CÔNG TY TNHH HẢI THÀNH")
    r_l.font.name = "Times New Roman"
    r_l.font.size = Pt(13)
    r_l.bold = True
    
    # Cột phải:
    # Dòng 1: CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM (Dấu huyền trên chữ A: HOÀ)
    # Dòng 2: Độc lập – Tự do – Hạnh phúc (Dấu gạch ngang dài en-dash)
    cell_right = tbl.cell(0, 1)
    p_qh = cell_right.paragraphs[0]
    p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_qh = p_qh.add_run("CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    r_qh.font.name = "Times New Roman"
    r_qh.font.size = Pt(12)
    r_qh.bold = True
    
    p_tn = cell_right.add_paragraph()
    p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tn = p_tn.add_run("Độc lập – Tự do – Hạnh phúc")
    r_tn.font.name = "Times New Roman"
    r_tn.font.size = Pt(13)
    r_tn.bold = True
    
    test_file = "storage/templates/test_user_screenshot.docx"
    doc.save(test_file)
    print(f"✔ Đã tạo file giả lập trường hợp lỗi của người dùng: {test_file}")
    
    validator = Decree30Validator(test_file, "cong_van_de_nghi")
    res = validator.validate_document()
    
    box1 = next((b for b in res["checklist"] if b["box_id"] == 1), None)
    print("\n--- KẾT QUẢ KIỂM THỬ Ô 1 VỚI CHỮ 'CỘNG HOÀ' ---")
    print("Passed:", box1["passed"])
    if not box1["passed"]:
        for err in box1["errors"]:
            print("Lỗi phát hiện:", err["issue"])
    else:
        print("🎉 THÀNH CÔNG: Ô 1 đã được nhận diện chính xác!")

if __name__ == "__main__":
    test_user_screenshot_case()
