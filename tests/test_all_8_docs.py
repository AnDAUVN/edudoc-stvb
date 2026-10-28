# -*- coding: utf-8 -*-
import os
import sys
import json

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from engine.validator import Decree30Validator

def set_page_margins(doc):
    for sec in doc.sections:
        sec.top_margin = Inches(0.79)     # 2.0 cm
        sec.bottom_margin = Inches(0.79)  # 2.0 cm
        sec.left_margin = Inches(1.18)    # 3.0 cm
        sec.right_margin = Inches(0.79)   # 2.0 cm

def create_header_table(doc, cq_top, cq_bottom, quoc_hieu, tieu_ngu, so_text, dia_danh_text, trich_yeu_5b=None):
    t_head = doc.add_table(rows=2, cols=2)
    t_head.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Cell (0,0): Cơ quan
    c00 = t_head.cell(0, 0)
    p_top = c00.paragraphs[0]
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_top = p_top.add_run(cq_top)
    r_top.font.name = "Times New Roman"
    r_top.font.size = Pt(12)

    p_bot = c00.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_bot = p_bot.add_run(cq_bottom)
    r_bot.font.name = "Times New Roman"
    r_bot.font.size = Pt(12)
    r_bot.bold = True

    # Cell (0,1): Quốc hiệu & Tiêu ngữ
    c01 = t_head.cell(0, 1)
    p_qh = c01.paragraphs[0]
    p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_qh = p_qh.add_run(quoc_hieu)
    r_qh.font.name = "Times New Roman"
    r_qh.font.size = Pt(12)
    r_qh.bold = True

    p_tn = c01.add_paragraph()
    p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tn = p_tn.add_run(tieu_ngu)
    r_tn.font.name = "Times New Roman"
    r_tn.font.size = Pt(13)
    r_tn.bold = True

    # Cell (1,0): Số ký hiệu (+ Ô 5b nếu là công văn)
    c10 = t_head.cell(1, 0)
    p_so = c10.paragraphs[0]
    p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_so.paragraph_format.space_after = Pt(6)
    r_so = p_so.add_run(so_text)
    r_so.font.name = "Times New Roman"
    r_so.font.size = Pt(13)

    if trich_yeu_5b:
        p_5b = c10.add_paragraph()
        p_5b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_5b.paragraph_format.space_before = Pt(6)
        r_5b = p_5b.add_run(trich_yeu_5b)
        r_5b.font.name = "Times New Roman"
        r_5b.font.size = Pt(12)

    # Cell (1,1): Địa danh ngày tháng
    c11 = t_head.cell(1, 1)
    p_dd = c11.paragraphs[0]
    p_dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dd = p_dd.add_run(dia_danh_text)
    r_dd.font.name = "Times New Roman"
    r_dd.font.size = Pt(13)
    r_dd.italic = True

def create_footer_table(doc, noi_nhan_list, chuc_vu_lines, ho_ten):
    t_foot = doc.add_table(rows=1, cols=2)
    t_foot.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Cell (0,0): Nơi nhận Ô 9b
    c0 = t_foot.cell(0, 0)
    p_nn = c0.paragraphs[0]
    r_nn = p_nn.add_run("Nơi nhận:")
    r_nn.font.name = "Times New Roman"
    r_nn.font.size = Pt(12)
    r_nn.bold = True
    r_nn.italic = True

    for item in noi_nhan_list:
        p_i = c0.add_paragraph()
        r_i = p_i.add_run(item)
        r_i.font.name = "Times New Roman"
        r_i.font.size = Pt(11)

    # Cell (0,1): Chức vụ, chữ ký, họ tên
    c1 = t_foot.cell(0, 1)
    for idx, cv in enumerate(chuc_vu_lines):
        p_cv = c1.paragraphs[0] if idx == 0 else c1.add_paragraph()
        p_cv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cv = p_cv.add_run(cv)
        r_cv.font.name = "Times New Roman"
        r_cv.font.size = Pt(13)
        r_cv.bold = True

    # 3-4 dòng trống ký tên
    for _ in range(3):
        p_empty = c1.add_paragraph()
        p_empty.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_ht = c1.add_paragraph()
    p_ht.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ht = p_ht.add_run(ho_ten)
    r_ht.font.name = "Times New Roman"
    r_ht.font.size = Pt(13)
    r_ht.bold = True

def create_sample_doc(doc_type_id, output_path):
    doc = Document()
    set_page_margins(doc)

    if doc_type_id == "bai_thu_hoach":
        create_header_table(
            doc,
            cq_top="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            cq_bottom="KHOA CÔNG NGHỆ THÔNG TIN",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 01/DNTT",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 25 tháng 10 năm 2026",
            trich_yeu_5b="V/v đề nghị hỗ trợ kinh phí"
        )
        # Ô 9a Kính gửi (Bắt buộc theo sơ đồ Công văn)
        p_kg = doc.add_paragraph()
        r_kg = p_kg.add_run("Kính gửi: Bộ Khoa học và Công nghệ.")
        r_kg.font.name = "Times New Roman"
        r_kg.font.size = Pt(13)

        # Ô 6 Thân bài
        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.line_spacing = 1.15
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Sau quá trình học tập và nghiên cứu các quy định của Nghị định 30/2020/NĐ-CP, người làm bài đã nắm vững các thành phần thể thức văn bản hành chính và thực hành soạn thảo thành thạo trên phần mềm soạn thảo.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        # Footer table
        create_footer_table(
            doc,
            noi_nhan_list=["- Như kính gửi;", "- Lưu: VT."],
            chuc_vu_lines=["HIỆU TRƯỞNG"],
            ho_ten="Nguyễn Văn A"
        )

    elif doc_type_id == "cong_van_de_nghi":
        create_header_table(
            doc,
            cq_top="BỘ GIÁO DỤC VÀ ĐÀO TẠO",
            cq_bottom="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 12/DNTT",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 15 tháng 01 năm 2026",
            trich_yeu_5b="V/v đề nghị hỗ trợ kinh phí"
        )
        p_kg = doc.add_paragraph()
        r_kg = p_kg.add_run("Kính gửi: Bộ Khoa học và Công nghệ.")
        r_kg.font.name = "Times New Roman"
        r_kg.font.size = Pt(13)

        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Trường Đại học Nguyễn Tất Thành xin gửi lời chào trân trọng đến Quý Bộ và trân trọng đề nghị xem xét hỗ trợ kinh phí nghiên cứu khoa học năm 2026.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        create_footer_table(
            doc,
            noi_nhan_list=["- Như kính gửi;", "- Lưu: VT."],
            chuc_vu_lines=["HIỆU TRƯỞNG"],
            ho_ten="Trần Ái Cầm"
        )

    elif doc_type_id == "thong_bao_nhiem_vu":
        create_header_table(
            doc,
            cq_top="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            cq_bottom="PHÒNG ĐÀO TẠO",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 45/TB-ĐT",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 02 tháng 02 năm 2026"
        )
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("THÔNG BÁO")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(13)
        r_t.bold = True

        p_y = doc.add_paragraph()
        p_y.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_y = p_y.add_run("Về việc phân công nhiệm vụ cán bộ coi thi kết thúc học phần học kỳ 1 năm học 2025 - 2026")
        r_y.font.name = "Times New Roman"
        r_y.font.size = Pt(13)
        r_y.bold = True

        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Thực hiện Kế hoạch tổ chức kỳ thi kết thúc học phần, Phòng Đào tạo thông báo đến toàn thể cán bộ, giảng viên có mặt đúng giờ để nhận nhiệm vụ coi thi.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        create_footer_table(
            doc,
            noi_nhan_list=["- Ban Giám hiệu (để báo cáo);", "- Cán bộ coi thi (để thực hiện);", "- Lưu: VT, ĐT."],
            chuc_vu_lines=["TL. HIỆU TRƯỞNG", "TRƯỞNG PHÒNG ĐÀO TẠO"],
            ho_ten="Nguyễn Văn An"
        )

    elif doc_type_id == "thong_bao_hoat_dong":
        create_header_table(
            doc,
            cq_top="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            cq_bottom="ĐOÀN THANH NIÊN",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 08/TB-ĐTN",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 08 tháng 03 năm 2026"
        )
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("THÔNG BÁO")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(13)
        r_t.bold = True

        p_y = doc.add_paragraph()
        p_y.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_y = p_y.add_run("Về việc tổ chức chiến dịch tình nguyện Mùa hè xanh năm 2026")
        r_y.font.name = "Times New Roman"
        r_y.font.size = Pt(13)
        r_y.bold = True

        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Đoàn Thanh niên Trường Đại học Nguyễn Tất Thành thông báo về việc phát động tuyển tình nguyện viên tham gia chiến dịch mùa hè xanh năm 2026.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        create_footer_table(
            doc,
            noi_nhan_list=["- Đảng ủy, Ban Giám hiệu (để báo cáo);", "- Các chi đoàn sinh viên;", "- Lưu: VT, ĐTN."],
            chuc_vu_lines=["BÍ THƯ ĐOÀN TRƯỜNG"],
            ho_ten="Lê Văn Bình"
        )

    elif doc_type_id == "to_trinh_xin_kinh_phi":
        create_header_table(
            doc,
            cq_top="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            cq_bottom="ĐOÀN THANH NIÊN",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 35/TTr-ĐTN",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 10 tháng 11 năm 2026"
        )
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("TỜ TRÌNH")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(13)
        r_t.bold = True

        p_y = doc.add_paragraph()
        p_y.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_y = p_y.add_run("Về việc xin cấp kinh phí tổ chức Hội thi Sáng tạo sinh viên NTTU năm 2026")
        r_y.font.name = "Times New Roman"
        r_y.font.size = Pt(13)
        r_y.bold = True

        p_kg = doc.add_paragraph()
        r_kg = p_kg.add_run("Kính gửi: Ban Giám hiệu Trường Đại học Nguyễn Tất Thành.")
        r_kg.font.name = "Times New Roman"
        r_kg.font.size = Pt(13)

        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Để chuẩn bị tốt cho Hội thi Sáng tạo sinh viên NTTU năm 2026, Đoàn trường kính trình Ban Giám hiệu xem xét phê duyệt kinh phí 90.000.000 đồng.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        create_footer_table(
            doc,
            noi_nhan_list=["- Như kính gửi;", "- Phòng Tài chính - Kế toán;", "- Lưu: VT, ĐTN."],
            chuc_vu_lines=["BÍ THƯ ĐOÀN TRƯỜNG"],
            ho_ten="Nguyễn Hoàng Long"
        )

    elif doc_type_id == "to_trinh_nhan_su":
        create_header_table(
            doc,
            cq_top="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            cq_bottom="KHOA CÔNG NGHỆ THÔNG TIN",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 15/TTr-CNTT",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 18 tháng 09 năm 2026"
        )
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("TỜ TRÌNH")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(13)
        r_t.bold = True

        p_y = doc.add_paragraph()
        p_y.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_y = p_y.add_run("Về việc xin bổ sung nhân sự giảng viên cơ hữu Bộ môn Kỹ thuật phần mềm")
        r_y.font.name = "Times New Roman"
        r_y.font.size = Pt(13)
        r_y.bold = True

        p_kg = doc.add_paragraph()
        r_kg = p_kg.add_run("Kính gửi: Ban Giám hiệu Trường Đại học Nguyễn Tất Thành.")
        r_kg.font.name = "Times New Roman"
        r_kg.font.size = Pt(13)

        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Căn cứ vào quy mô phát triển sinh viên ngành Công nghệ thông tin, Khoa kính trình Ban Giám hiệu phê duyệt tuyển dụng bổ sung 03 giảng viên cơ hữu.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        create_footer_table(
            doc,
            noi_nhan_list=["- Như kính gửi;", "- Phòng Tổ chức - Nhân sự;", "- Lưu: VT, Khoa CNTT."],
            chuc_vu_lines=["TRƯỞNG KHOA"],
            ho_ten="Hoàng Minh Trí"
        )

    elif doc_type_id == "to_trinh_thiet_bi":
        create_header_table(
            doc,
            cq_top="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            cq_bottom="TRUNG TÂM ĐÀO TẠO KỸ NĂNG NGHỀ",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 22/TTr-KNN",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 24 tháng 10 năm 2026"
        )
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("TỜ TRÌNH")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(13)
        r_t.bold = True

        p_y = doc.add_paragraph()
        p_y.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_y = p_y.add_run("Về việc xin trang bị bổ sung máy chiếu và dàn máy vi tính phòng thực hành")
        r_y.font.name = "Times New Roman"
        r_y.font.size = Pt(13)
        r_y.bold = True

        p_kg = doc.add_paragraph()
        r_kg = p_kg.add_run("Kính gửi: Ban Giám hiệu Trường Đại học Nguyễn Tất Thành.")
        r_kg.font.name = "Times New Roman"
        r_kg.font.size = Pt(13)

        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Nhằm phục vụ tốt kỳ thi đánh giá kỹ năng nghề bậc 3 cho sinh viên, Trung tâm kính đề nghị Ban Giám hiệu phê duyệt mua sắm 05 máy chiếu và 20 bộ máy tính.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        create_footer_table(
            doc,
            noi_nhan_list=["- Như kính gửi;", "- Phòng Quản trị thiết bị;", "- Lưu: VT, TT KNN."],
            chuc_vu_lines=["KT. GIÁM ĐỐC", "PHÓ GIÁM ĐỐC"],
            ho_ten="Võ Thị Kim Anh"
        )

    elif doc_type_id == "bao_cao":
        create_header_table(
            doc,
            cq_top="TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH",
            cq_bottom="KHOA CÔNG NGHỆ THÔNG TIN",
            quoc_hieu="CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            tieu_ngu="Độc lập - Tự do - Hạnh phúc",
            so_text="Số: 18/BC-CNTT",
            dia_danh_text="Thành phố Hồ Chí Minh, ngày 25 tháng 10 năm 2026"
        )
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_t.add_run("BÁO CÁO")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(13)
        r_t.bold = True

        p_y = doc.add_paragraph()
        p_y.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_y = p_y.add_run("Về tình hình thực hiện công tác đào tạo và nghiên cứu khoa học học kỳ 1 năm học 2025 - 2026")
        r_y.font.name = "Times New Roman"
        r_y.font.size = Pt(13)
        r_y.bold = True

        p_kg = doc.add_paragraph()
        r_kg = p_kg.add_run("Kính gửi: Ban Giám hiệu Trường Đại học Nguyễn Tất Thành.")
        r_kg.font.name = "Times New Roman"
        r_kg.font.size = Pt(13)

        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(0.5)
        p_b.paragraph_format.space_after = Pt(6)
        r_b = p_b.add_run("Thực hiện Kế hoạch năm học 2025 - 2026 của Nhà trường, Khoa Công nghệ thông tin kính báo cáo tình hình hoàn thành 100% khối lượng giờ giảng và nghiệm thu đúng hạn 05 đề tài nghiên cứu cấp cơ sở.")
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(13)

        create_footer_table(
            doc,
            noi_nhan_list=["- Như kính gửi;", "- Phòng Đào tạo;", "- Lưu: VT, Khoa CNTT."],
            chuc_vu_lines=["TRƯỞNG KHOA"],
            ho_ten="Hoàng Minh Trí"
        )

    doc.save(output_path)
    return output_path

def test_all_8_documents():
    print("=" * 70)
    print("  KIỂM TRA CHỨC NĂNG & THẨM ĐỊNH THỂ THỨC CHO 8 LOẠI VĂN BẢN")
    print("=" * 70)

    # 1. Check templates configuration
    templates_path = os.path.join(BASE_DIR, "data", "document_templates.json")
    with open(templates_path, "r", encoding="utf-8") as f:
        templates = json.load(f)

    expected_keys = [
        "bai_thu_hoach",
        "cong_van_de_nghi",
        "thong_bao_nhiem_vu",
        "thong_bao_hoat_dong",
        "to_trinh_xin_kinh_phi",
        "to_trinh_nhan_su",
        "to_trinh_thiet_bi",
        "bao_cao"
    ]

    actual_keys = list(templates.keys())
    assert actual_keys == expected_keys, f"Thứ tự văn bản không khớp!\nMong muốn: {expected_keys}\nThực tế: {actual_keys}"
    print(" [+] Thứ tự và danh sách 8 loại văn bản trong templates hoàn toàn khớp!")

    # 2. Check quiz bank configuration
    questions_path = os.path.join(BASE_DIR, "data", "questions.json")
    with open(questions_path, "r", encoding="utf-8") as f:
        questions = json.load(f)

    for doc_id in expected_keys:
        assert doc_id in questions, f"Thiếu câu hỏi cho {doc_id}"
        q_cnt = len(questions[doc_id])
        assert q_cnt >= 5, f"{doc_id} có ít hơn 5 câu hỏi ({q_cnt})"
        print(f" [+] Ngân hàng câu hỏi {doc_id}: {q_cnt} câu hỏi.")

    # 3. Create and Validate Word files for all 8 types
    test_out_dir = os.path.join(BASE_DIR, "storage", "test_out")
    os.makedirs(test_out_dir, exist_ok=True)

    print("\n--- CHẠY THẨM ĐỊNH NGHỊ ĐỊNH 30 CHO CẢ 8 LOẠI VĂN BẢN ---")
    for doc_id in expected_keys:
        out_file = os.path.join(test_out_dir, f"test_perfect_{doc_id}.docx")
        create_sample_doc(doc_id, out_file)
        
        validator = Decree30Validator(out_file, doc_id, template_info=templates[doc_id])
        result = validator.validate_document()

        passed_count = sum(1 for item in result["checklist"] if item["passed"])
        total_count = len(result["checklist"])
        score = result["percentage"]
        
        errors = []
        for box in result["checklist"]:
            if not box["passed"]:
                for err in box.get("errors", []):
                    errors.append(f"[Ô {box.get('box_id')}]: {err.get('issue')}")

        print(f" ✔ {doc_id:22s}: Điểm {score:5.1f}/100 | Check: {passed_count}/{total_count} passed | Profile: {result['font_profile']}")
        if len(errors) > 0:
            print(f"   Lỗi phát hiện ({len(errors)}):")
            for err in errors:
                print(f"     - {err}")
        
        assert score == 100.0, f"Văn bản {doc_id} chưa đạt 100 điểm ({score})!"
        assert len(errors) == 0, f"Văn bản {doc_id} còn lỗi!"

    print("\n" + "=" * 70)
    print(" 🎉 TOÀN BỘ 8 LOẠI VĂN BẢN ĐẠT 100% TIÊU CHUẨN NGHỊ ĐỊNH 30/2020/NĐ-CP!")
    print("=" * 70)

if __name__ == "__main__":
    test_all_8_documents()
