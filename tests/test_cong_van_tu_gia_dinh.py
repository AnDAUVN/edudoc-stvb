# -*- coding: utf-8 -*-
"""
Kiểm thử chuyên sâu: Bài tập Công văn đề nghị - Tự giả định tất cả các mục
Tuân thủ Sơ đồ bố trí Công văn (Hình 2 theo Nghị định 30/2020/NĐ-CP)
"""
import os
import sys
import unittest
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from app import app, SESSION_STATE_FILE, save_json_file, load_json_file
from engine.validator import Decree30Validator

class TestCongVanTuGiaDinh(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()
        self.temp_dir = os.path.join(BASE_DIR, "storage", "test_temp")
        os.makedirs(self.temp_dir, exist_ok=True)
        self.orig_state = load_json_file(SESSION_STATE_FILE, {})

    def tearDown(self):
        save_json_file(SESSION_STATE_FILE, self.orig_state)

    def test_01_default_active_doc_is_cong_van_de_nghi(self):
        """Mặc định hệ thống giao bài tập Công văn đề nghị"""
        save_json_file(SESSION_STATE_FILE, {
            "is_open": True,
            "active_doc_type": "cong_van_de_nghi",
            "enforce_roster": False
        })
        res = self.client.get("/api/session-state")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.get_json().get("active_doc_type"), "cong_van_de_nghi")

        res_student = self.client.get("/")
        self.assertEqual(res_student.status_code, 200)
        html = res_student.get_data(as_text=True)
        self.assertIn('value="cong_van_de_nghi" selected', html)

    def test_02_cong_van_tu_gia_dinh_perfect(self):
        """Văn bản Công văn đề nghị tự giả định toàn bộ nội dung đạt 100% điểm thể thức Hình 2 NĐ 30"""
        doc_path = os.path.join(self.temp_dir, "test_cv_tu_gia_dinh_perfect.docx")
        doc = Document()

        # Căn lề chuẩn
        for section in doc.sections:
            section.top_margin = Inches(0.85)     # ~2.16cm
            section.bottom_margin = Inches(0.85)  # ~2.16cm
            section.left_margin = Inches(1.2)     # ~3.05cm
            section.right_margin = Inches(0.7)    # ~1.78cm

        # BẢNG ĐẦU TRANG 2x2
        ht = doc.add_table(rows=2, cols=2)
        ht.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Hàng 1 Cột 1: Ô 2 Cơ quan (Tự giả định)
        c00 = ht.cell(0, 0)
        p_cq_top = c00.paragraphs[0]
        p_cq_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p_cq_top.add_run("UBND TỈNH ĐỒNG NAI")
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(12)
        r1.bold = False

        p_cq_bot = c00.add_paragraph()
        p_cq_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p_cq_bot.add_run("SỞ TƯ PHÁP")
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(12)
        r2.bold = True

        # Hàng 1 Cột 2: Ô 1 Quốc hiệu & Tiêu ngữ
        c01 = ht.cell(0, 1)
        p_qh = c01.paragraphs[0]
        p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rqh = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
        rqh.font.name = "Times New Roman"
        rqh.font.size = Pt(12)
        rqh.bold = True

        p_tn = c01.add_paragraph()
        p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rtn = p_tn.add_run("Độc lập - Tự do - Hạnh phúc")
        rtn.font.name = "Times New Roman"
        rtn.font.size = Pt(13)
        rtn.bold = True

        # Hàng 2 Cột 1: Ô 3 Số ký hiệu & Ô 5b Trích yếu V/v (Tự giả định hoàn toàn)
        c10 = ht.cell(1, 0)
        p_so = c10.paragraphs[0]
        p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rso = p_so.add_run("Số: 99/STP-VP")  # Tự giả định, không có CV, không ép MSSV
        rso.font.name = "Times New Roman"
        rso.font.size = Pt(13)
        rso.bold = False

        p_5b = c10.add_paragraph()
        p_5b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_5b.paragraph_format.space_before = Pt(6)
        r5b = p_5b.add_run("V/v phối hợp tổ chức tuyên truyền phổ biến giáo dục pháp luật năm 2026")
        r5b.font.name = "Times New Roman"
        r5b.font.size = Pt(12)
        r5b.bold = False
        r5b.italic = False

        # Hàng 2 Cột 2: Ô 4 Địa danh ngày tháng
        c11 = ht.cell(1, 1)
        p_dd = c11.paragraphs[0]
        p_dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rdd = p_dd.add_run("Đồng Nai, ngày 15 tháng 03 năm 2026")
        rdd.font.name = "Times New Roman"
        rdd.font.size = Pt(13)
        rdd.italic = True
        rdd.bold = False

        # NGOÀI BẢNG: Ô 9a (Kính gửi) - BẮT BUỘC CÓ TRƯỚC Ô 6
        p_kg = doc.add_paragraph()
        p_kg.paragraph_format.left_indent = Inches(0.5)
        rkg = p_kg.add_run("Kính gửi: Ủy ban nhân dân các huyện, thành phố.")
        rkg.font.name = "Times New Roman"
        rkg.font.size = Pt(13)
        rkg.bold = False

        # Ô 6: NỘI DUNG VĂN BẢN (Tự giả định)
        p_body1 = doc.add_paragraph()
        p_body1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body1.paragraph_format.first_line_indent = Inches(0.45)
        p_body1.paragraph_format.space_after = Pt(6)
        p_body1.paragraph_format.line_spacing = 1.15
        rb1 = p_body1.add_run("Thực hiện Kế hoạch công tác phổ biến giáo dục pháp luật năm 2026 của Ủy ban nhân dân tỉnh, Sở Tư pháp đề nghị các đơn vị phối hợp triển khai các chuyên đề tuyên truyền trọng tâm tại địa bàn cơ sở.")
        rb1.font.name = "Times New Roman"
        rb1.font.size = Pt(13)

        p_body2 = doc.add_paragraph()
        p_body2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body2.paragraph_format.first_line_indent = Inches(0.45)
        p_body2.paragraph_format.space_after = Pt(6)
        p_body2.paragraph_format.line_spacing = 1.15
        rb2 = p_body2.add_run("Đề nghị Ủy ban nhân dân các huyện, thành phố quan tâm chỉ đạo Phòng Tư pháp và các đơn vị liên quan lập danh sách cán bộ tham gia tập huấn trước ngày 30 tháng 03 năm 2026.")
        rb2.font.name = "Times New Roman"
        rb2.font.size = Pt(13)

        # BẢNG CUỐI TRANG 1x2
        ft = doc.add_table(rows=1, cols=2)
        ft.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Cột trái: Ô 9b Nơi nhận
        fc0 = ft.cell(0, 0)
        p_nn_head = fc0.paragraphs[0]
        rnn_h = p_nn_head.add_run("Nơi nhận:")
        rnn_h.font.name = "Times New Roman"
        rnn_h.font.size = Pt(12)
        rnn_h.bold = True
        rnn_h.italic = True

        p_nn1 = fc0.add_paragraph()
        rnn1 = p_nn1.add_run("- Như kính gửi;")
        rnn1.font.name = "Times New Roman"
        rnn1.font.size = Pt(11)

        p_nn2 = fc0.add_paragraph()
        rnn2 = p_nn2.add_run("- Lưu: VT, VP.")
        rnn2.font.name = "Times New Roman"
        rnn2.font.size = Pt(11)

        # Cột phải: Ô 7 & Ô 8 Chức vụ & Họ tên (Tự giả định)
        fc1 = ft.cell(0, 1)
        p_cv = fc1.paragraphs[0]
        p_cv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rcv = p_cv.add_run("GIÁM ĐỐC")
        rcv.font.name = "Times New Roman"
        rcv.font.size = Pt(13)
        rcv.bold = True

        for _ in range(3):
            fc1.add_paragraph()

        p_name = fc1.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rname = p_name.add_run("Nguyễn Văn Hùng")
        rname.font.name = "Times New Roman"
        rname.font.size = Pt(13)
        rname.bold = True

        doc.save(doc_path)

        # Kiểm định với MSSV khác số 99 để xác nhận không bị ép số MSSV
        validator = Decree30Validator(doc_path, "cong_van_de_nghi", mssv="2100001234")
        report = validator.validate_document()

        failed_items = [b for b in report["checklist"] if not b["passed"]]
        self.assertEqual(len(failed_items), 0, f"Công văn tự giả định chuẩn phải 0 lỗi! Các mục lỗi: {failed_items}")
        self.assertTrue(report["can_submit"])
        self.assertEqual(report["percentage"], 100.0)
        self.assertEqual(report["passed_count"], report["total_count"])

    def test_03_cong_van_sai_so_do_co_5a_bi_bat_loi(self):
        """Nếu sinh viên chèn Ô 5a 'CÔNG VĂN ĐỀ NGHỊ' vào thân bài, validator bắt lỗi sơ đồ Hình 2"""
        doc_path = os.path.join(self.temp_dir, "test_cv_co_5a_sai.docx")
        doc = Document()
        ht = doc.add_table(rows=2, cols=2)
        # Điền tạm Ô 5a vi phạm
        p_sai = doc.add_paragraph("CÔNG VĂN ĐỀ NGHỊ")
        p_sai.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("Kính gửi: Sở Giáo dục và Đào tạo")
        doc.add_paragraph("Nội dung bài viết...")
        ft = doc.add_table(rows=1, cols=2)
        doc.save(doc_path)

        validator = Decree30Validator(doc_path, "cong_van_de_nghi", mssv="2100001234")
        report = validator.validate_document()
        has_5a_err = any(
            any("KHÔNG CÓ TÊN LOẠI (Ô 5a)" in e["issue"] or "Bỏ (xóa) dòng chữ này ra khỏi sơ đồ văn bản" in e.get("fix", "") for e in box["errors"])
            for box in report["checklist"]
        )
        self.assertTrue(has_5a_err, "Phải bắt lỗi sai sơ đồ bố trí: Công văn không có Ô 5a!")

    def test_04_cong_van_co_tu_thong_bao_trong_noi_dung_khong_bi_nham_la_5a(self):
        """Công văn có các từ ngữ 'thông báo', 'báo cáo' trong phần nội dung thân bài (Ô 6) không được nhận nhầm là Ô 5a"""
        doc_path = os.path.join(self.temp_dir, "test_cv_body_with_kw.docx")
        doc = Document()
        ht = doc.add_table(rows=2, cols=2)
        # Bảng đầu trang đúng chuẩn
        ht.cell(0, 0).paragraphs[0].text = "CÔNG TY ÁNH DƯƠNG"
        ht.cell(0, 1).paragraphs[0].text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
        ht.cell(1, 0).paragraphs[0].text = "Số: 122/AD-HC"
        p_vv = ht.cell(1, 0).add_paragraph("V/v khắc phục sự cố")
        ht.cell(1, 1).paragraphs[0].text = "Hà Nội, ngày 06 tháng 08 năm 2026"

        # Kính gửi Ô 9a
        doc.add_paragraph("Kính gửi: Sở Công Thương quận Gò Vấp;")
        doc.add_paragraph("Công ty Điện Lực Gò Vấp.")

        # Thân bài Ô 6 chứa các từ thông báo, báo cáo
        doc.add_paragraph("Hiện nay Công ty Ánh Dương đang vào giai đoạn cao điểm.")
        doc.add_paragraph("Chấm dứt ngay tình trạng cắt điện đột ngột không thông báo trước.")
        doc.add_paragraph("Thực hiện nghiêm túc quy trình thông báo ngừng, giảm cung cấp điện.")
        
        ft = doc.add_table(rows=1, cols=2)
        ft.cell(0, 0).paragraphs[0].text = "Nơi nhận:\n- Như trên;\n- Lưu: VT."
        ft.cell(0, 1).paragraphs[0].text = "GIÁM ĐỐC\n\n\nTrần Thị Hiền"
        doc.save(doc_path)

        validator = Decree30Validator(doc_path, "cong_van_de_nghi", mssv="2100001234")
        report = validator.validate_document()
        has_false_5a_err = any(
            any("KHÔNG CÓ TÊN LOẠI (Ô 5a)" in e.get("issue", "") for e in box.get("errors", []))
            for box in report.get("checklist", [])
        )
        self.assertFalse(has_false_5a_err, "Công văn có từ 'thông báo' trong câu thân bài KHÔNG ĐƯỢC báo lỗi Ô 5a!")

    def test_05_bullet_points_in_body_do_not_trigger_first_line_indent_error(self):
        """Các dòng gạch đầu dòng / bullet point / danh sách liệt kê trong thân bài (Ô 6) không bị phạt lỗi thụt lề 1-1.27cm"""
        doc_path = os.path.join(self.temp_dir, "test_cv_bullets_in_body.docx")
        doc = Document()
        ht = doc.add_table(rows=2, cols=2)
        ht.cell(0, 0).paragraphs[0].text = "CÔNG TY ÁNH DƯƠNG"
        ht.cell(0, 1).paragraphs[0].text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
        ht.cell(1, 0).paragraphs[0].text = "Số: 122/AD-HC"
        ht.cell(1, 0).add_paragraph("V/v giải quyết kiến nghị mất điện")
        ht.cell(1, 1).paragraphs[0].text = "Hà Nội, ngày 06 tháng 08 năm 2026"

        # Ô 9a
        doc.add_paragraph("Kính gửi: Sở Công Thương quận Gò Vấp;")
        doc.add_paragraph("Công ty Điện Lực Gò Vấp.")

        # Ô 6: Thân bài có đoạn văn thụt lề và các bullet points
        p_narrative = doc.add_paragraph()
        p_narrative.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_narrative.paragraph_format.first_line_indent = Inches(0.45)
        p_narrative.add_run("Phòng Hành chính đề nghị Sở Công Thương quận Gò Vấp và Công ty Điện Lực Gò Vấp làm rõ nguyên nhân và tình trạng cúp điện như sau:")

        # Bullet 1 (dùng ký tự •)
        p_bullet1 = doc.add_paragraph()
        p_bullet1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_bullet1.paragraph_format.first_line_indent = None  # Bullet trong Word không có first line indent
        p_bullet1.add_run("• Chấm dứt ngay tình trạng cắt điện đột ngột không thông báo trước cho các đơn vị kinh doanh.")

        # Bullet 2 (dùng gạch đầu dòng -)
        p_bullet2 = doc.add_paragraph()
        p_bullet2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_bullet2.paragraph_format.first_line_indent = None
        p_bullet2.add_run("- Thực hiện nghiêm túc quy trình thông báo ngừng, giảm cung cấp điện qua các kênh liên lạc theo đúng luật.")

        # Bullet 3 (dùng số 1.)
        p_bullet3 = doc.add_paragraph()
        p_bullet3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_bullet3.paragraph_format.first_line_indent = None
        p_bullet3.add_run("1. Kiểm tra, nâng cấp hệ thống hạ tầng kỹ thuật nhằm khắc phục triệt để các sự cố sụt áp trên địa bàn.")

        p_close = doc.add_paragraph()
        p_close.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_close.paragraph_format.first_line_indent = Inches(0.45)
        p_close.add_run("Kính mong nhận được sự quan tâm và giúp đỡ của Quý cơ quan để vấn đề trên sớm được giải quyết thỏa đáng.")

        ft = doc.add_table(rows=1, cols=2)
        ft.cell(0, 0).paragraphs[0].text = "Nơi nhận:\n- Như trên;\n- Lưu: VT."
        ft.cell(0, 1).paragraphs[0].text = "GIÁM ĐỐC\n\n\nTrần Thị Hiền"
        doc.save(doc_path)

        validator = Decree30Validator(doc_path, "cong_van_de_nghi", mssv="2100001234")
        report = validator.validate_document()
        box6 = next((b for b in report.get("checklist", []) if str(b.get("box_id")) == "6"), None)
        self.assertIsNotNone(box6, "Phải tìm thấy kết quả kiểm tra Ô 6")
        
        # 1. Các dòng bullet chưa lùi vào phải bị báo lỗi
        has_list_indent_err = any("chưa lùi vào so với ý chính" in e.get("issue", "") for e in box6.get("errors", []))
        self.assertTrue(has_list_indent_err, "Phải bắt lỗi nếu dòng bullet/danh sách liệt kê chưa lùi vào so với ý chính!")

        # 2. Khi lùi vào (bằng Tab hoặc left_indent hoặc Word bullets), không còn bị báo lỗi
        p_bullet1.paragraph_format.left_indent = Inches(0.5)
        p_bullet2.text = "\t" + p_bullet2.text
        p_bullet3.paragraph_format.first_line_indent = Inches(0.5)
        doc_indented_path = os.path.join(self.temp_dir, "test_cv_bullets_indented_ok.docx")
        doc.save(doc_indented_path)

        v_ok = Decree30Validator(doc_indented_path, "cong_van_de_nghi", mssv="2100001234")
        report_ok = v_ok.validate_document()
        box6_ok = next((b for b in report_ok.get("checklist", []) if str(b.get("box_id")) == "6"), None)
        list_errors_ok = [e for e in box6_ok.get("errors", []) if "chưa lùi vào so với ý chính" in e.get("issue", "")]
        self.assertEqual(len(list_errors_ok), 0, f"Dòng danh sách đã lùi vào không được báo lỗi! Lỗi: {list_errors_ok}")
        indent_errors_ok = [e for e in box6_ok.get("errors", []) if "thụt vào 1 - 1.27 cm" in e.get("issue", "")]
        self.assertEqual(len(indent_errors_ok), 0, f"Đoạn văn chính chuẩn và bullet đã lùi vào không được báo lỗi thụt dòng!")

    def test_06_punctuation_spacing_luu_vt(self):
        """Bắt lỗi quy tắc khoảng cách trống khi viết dính liền sau dấu hai chấm: '- Lưu:VT.' thay vì '- Lưu: VT.'"""
        doc_path = os.path.join(self.temp_dir, "test_cv_luu_vt_no_space.docx")
        doc = Document()
        ht = doc.add_table(rows=2, cols=2)
        doc.add_paragraph("Kính gửi: Sở Công Thương quận Gò Vấp.")
        doc.add_paragraph("Nội dung bài viết...")
        ft = doc.add_table(rows=1, cols=2)
        # Ô 9b cố tình viết dính liền Lưu:VT.
        p_nn_title = ft.cell(0, 0).paragraphs[0]
        p_nn_title.text = "Nơi nhận:"
        p_nn_title.runs[0].italic = True
        p_nn_title.runs[0].bold = True
        p_nn_title.runs[0].font.size = Pt(12)

        p_nn_item = ft.cell(0, 0).add_paragraph("- Lưu:VT.")
        p_nn_item.runs[0].font.size = Pt(11)

        p_cv = ft.cell(0, 1).paragraphs[0]
        p_cv.text = "GIÁM ĐỐC"
        p_cv.runs[0].bold = True
        p_cv.runs[0].font.size = Pt(13)
        for _ in range(3):
            ft.cell(0, 1).add_paragraph()
        p_name = ft.cell(0, 1).add_paragraph("Trần Thị Hiền")
        p_name.runs[0].bold = True
        p_name.runs[0].font.size = Pt(13)
        doc.save(doc_path)

        validator = Decree30Validator(doc_path, "cong_van_de_nghi", mssv="2100001234")
        report = validator.validate_document()
        box9b = next((b for b in report.get("checklist", []) if str(b.get("box_id")) == "9b"), None)
        self.assertIsNotNone(box9b, "Phải có Ô 9b")
        has_spacing_err = any("Dấu hai chấm" in e.get("issue", "") and "Lưu:VT" in e.get("issue", "") for e in box9b.get("errors", []))
        self.assertTrue(has_spacing_err, "Phải bắt lỗi khoảng cách trống: Thiếu khoảng trắng sau dấu hai chấm trong 'Lưu:VT.'!")

        # Khi sửa thành '- Lưu: VT.' -> Hết lỗi dấu hai chấm
        p_nn_item.text = "- Lưu: VT."
        p_nn_item.runs[0].font.size = Pt(11)
        doc.save(doc_path)
        v_fixed = Decree30Validator(doc_path, "cong_van_de_nghi", mssv="2100001234")
        rep_fixed = v_fixed.validate_document()
        box9b_fixed = next((b for b in rep_fixed.get("checklist", []) if str(b.get("box_id")) == "9b"), None)
        has_spacing_err_fixed = any("Dấu hai chấm" in e.get("issue", "") and "Lưu:VT" in e.get("issue", "") for e in box9b_fixed.get("errors", []))
        self.assertFalse(has_spacing_err_fixed, "Sau khi sửa thành '- Lưu: VT.' không còn bị báo lỗi dấu hai chấm!")

if __name__ == "__main__":
    unittest.main()
