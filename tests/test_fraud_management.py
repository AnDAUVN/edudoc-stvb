import requests
import json
import os

BASE_URL = "http://127.0.0.1:5000"

def test_fraud_management():
    print("=== KIỂM THỬ QUẢN LÝ NHẬT KÝ VI PHẠM & CÔ LẬP THEO MÃ LỚP ===")
    
    # 1. Kiểm tra API lấy dữ liệu giảng viên
    res = requests.get(f"{BASE_URL}/api/teacher/data")
    assert res.status_code == 200, f"Failed /api/teacher/data: {res.text}"
    data = res.json()
    assert "fraud_logs" in data, "Thiếu trường fraud_logs trong teacher/data"
    assert "classes" in data, "Thiếu trường classes trong teacher/data"
    print(" [+] API /api/teacher/data hoạt động tốt.")

    # 2. Xóa sạch dữ liệu vi phạm ban đầu để test
    res = requests.post(f"{BASE_URL}/api/teacher/fraud/clear", json={"class_code": "ALL"})
    assert res.status_code == 200
    res = requests.get(f"{BASE_URL}/api/teacher/data")
    assert len(res.json()["fraud_logs"]) == 0
    print(" [+] API /api/teacher/fraud/clear (ALL) hoạt động tốt, đã làm sạch nhật ký.")

    # 3. Giả lập một sinh viên nộp file không hợp lệ thuộc lớp KN01
    requests.post(f"{BASE_URL}/api/teacher/students/add", json={
        "class_code": "KN01",
        "sbd": "01",
        "mssv": "2280600111",
        "name": "Nguyễn Văn Gian Lận"
    })
    # Tạo một file docx rỗng bất kỳ
    from docx import Document
    test_docx = "temp_unauthorized.docx"
    doc = Document()
    doc.add_paragraph("Nội dung tùy tiện bên ngoài")
    doc.save(test_docx)

    with open(test_docx, "rb") as f:
        files = {"file": (test_docx, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        formData = {
            "class_code": "KN01",
            "sbd": "01",
            "mssv": "2280600111",
            "name": "Nguyễn Văn Gian Lận",
            "doc_type": "cong_van_de_nghi"
        }
        res = requests.post(f"{BASE_URL}/api/validate", files=files, data=formData)
    
    if os.path.exists(test_docx):
        os.remove(test_docx)

    assert res.status_code == 200
    val_data = res.json()
    assert val_data.get("status") == "FRAUD_DETECTED", f"Mong muốn FRAUD_DETECTED, nhận: {val_data}"
    print(" [+] Hệ thống đã bắt thành công gian lận FILE_KHONG_HOP_LE.")

    # 4. Kiểm tra xem fraud_logs đã lưu đúng mã lớp KN01 và có ID chưa
    res = requests.get(f"{BASE_URL}/api/teacher/data")
    logs = res.json()["fraud_logs"]
    assert len(logs) == 1, f"Mong muốn có 1 log vi phạm, có: {len(logs)}"
    log_item = logs[0]
    assert log_item.get("class_code") == "KN01", f"Mã lớp không đúng: {log_item.get('class_code')}"
    assert "id" in log_item and len(log_item["id"]) > 0, "Log vi phạm thiếu trường 'id'"
    assert log_item["fraud_data"]["student_a"]["name"] == "Nguyễn Văn Gian Lận"
    assert log_item["fraud_data"]["student_a"]["class_code"] == "KN01"
    log_id = log_item["id"]
    print(f" [+] Hồ sơ vi phạm đã lưu đúng mã lớp 'KN01' và ID: '{log_id}'.")

    # 5. Giả lập một sinh viên khác vi phạm ở lớp KN02
    requests.post(f"{BASE_URL}/api/teacher/students/add", json={
        "class_code": "KN02",
        "sbd": "02",
        "mssv": "2280600222",
        "name": "Trần Thị B"
    })
    doc = Document()
    doc.add_paragraph("Nội dung bài làm sao chép lớp khác")
    doc.save(test_docx)

    with open(test_docx, "rb") as f:
        files = {"file": (test_docx, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        formData = {
            "class_code": "KN02",
            "sbd": "02",
            "mssv": "2280600222",
            "name": "Trần Thị B",
            "doc_type": "thong_bao_nhiem_vu"
        }
        res = requests.post(f"{BASE_URL}/api/validate", files=files, data=formData)
    
    if os.path.exists(test_docx):
        os.remove(test_docx)

    res = requests.get(f"{BASE_URL}/api/teacher/data")
    logs = res.json()["fraud_logs"]
    assert len(logs) == 2, f"Mong muốn có 2 log vi phạm, có: {len(logs)}"
    print(" [+] Đã ghi nhận thêm 1 vi phạm ở lớp KN02 (Tổng 2 vi phạm: 1 KN01, 1 KN02).")

    # 6. Kiểm tra xóa đơn lẻ 1 log theo ID (xóa sinh viên lớp KN01)
    res = requests.post(f"{BASE_URL}/api/teacher/fraud/delete", json={"id": log_id})
    assert res.status_code == 200
    res = requests.get(f"{BASE_URL}/api/teacher/data")
    logs = res.json()["fraud_logs"]
    assert len(logs) == 1, f"Mong muốn còn 1 log sau khi xóa log_id, có: {len(logs)}"
    assert logs[0]["class_code"] == "KN02", "Log còn lại phải là của lớp KN02"
    print(" [+] API /api/teacher/fraud/delete xóa đơn lẻ theo ID hoạt động chính xác 100%.")

    # 7. Kiểm tra xóa theo mã lớp (xóa lớp KN02)
    res = requests.post(f"{BASE_URL}/api/teacher/fraud/clear", json={"class_code": "KN02"})
    assert res.status_code == 200
    res = requests.get(f"{BASE_URL}/api/teacher/data")
    logs = res.json()["fraud_logs"]
    assert len(logs) == 0, f"Mong muốn còn 0 log sau khi xóa lớp KN02, có: {len(logs)}"
    print(" [+] API /api/teacher/fraud/clear theo mã lớp KN02 hoạt động chính xác 100%.")

    print("\n>>> TẤT CẢ CÁC BÀI TEST QUẢN LÝ VI PHẠM & CÔ LẬP THEO MÃ LỚP ĐẠT 100%! <<<")

def ensure_server():
    import socket, subprocess, sys, time
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.settimeout(0.5)
        if s.connect_ex(("127.0.0.1", 5000)) == 0:
            return None
    root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    proc = subprocess.Popen([sys.executable, "app.py"], cwd=root, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for _ in range(20):
        time.sleep(0.5)
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.settimeout(0.5)
            if s.connect_ex(("127.0.0.1", 5000)) == 0:
                break
    return proc

if __name__ == "__main__":
    proc = ensure_server()
    try:
        test_fraud_management()
    finally:
        if proc:
            proc.terminate()
