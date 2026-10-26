import os
import sys
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Thêm thư mục hiện tại vào sys.path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from engine.security import (
    create_student_raw_template,
    check_anti_cheat,
    encode_hidden_mssv_to_text,
    extract_all_hidden_mssvs
)
from engine.validator import Decree30Validator

def test_full_pipeline():
    print("=" * 60)
    print("BẮT ĐẦU KIỂM THỬ HỆ THỐNG EDUDOC - STVB NGHỊ ĐỊNH 30")
    print("=" * 60)

    # 1. ĐỌC DỮ LIỆU ĐỀ BÀI
    with open("data/document_templates.json", "r", encoding="utf-8") as f:
        templates = json.load(f)
    tmpl = templates["thong_bao_nhiem_vu"]

    # 2. TẠO FILE CHO SINH VIÊN B (Người bị chép bài)
    file_b = "storage/templates/18_2100009876_TRANTHIB_ThongBaoNhiemVu.docx"
    create_student_raw_template(
        sbd="18",
        mssv="2100009876",
        student_name="Trần Thị B",
        doc_type_id="thong_bao_nhiem_vu",
        raw_text=tmpl["raw_text"],
        output_path=file_b
    )
    print(f"✔ Đã tạo file mẫu thô cho Sinh viên B: {file_b}")

    # 3. TẠO FILE CHO SINH VIÊN A (Người nộp)
    file_a = "storage/templates/01_2100001234_NGUYENVANA_ThongBaoNhiemVu.docx"
    create_student_raw_template(
        sbd="01",
        mssv="2100001234",
        student_name="Nguyễn Văn A",
        doc_type_id="thong_bao_nhiem_vu",
        raw_text=tmpl["raw_text"],
        output_path=file_a
    )
    print(f"✔ Đã tạo file mẫu thô cho Sinh viên A: {file_a}")

    # 4. TEST TÌNH HUỐNG GIAN LẬN 1: SV A nộp nguyên file của SV B (Đổi tên file)
    print("\n--- TEST GIAN LẬN 1: SV A NỘP NGUYÊN FILE CỦA SV B ---")
    fraud_res_1 = check_anti_cheat(
        docx_path=file_b, # A lấy file của B nộp
        current_sbd="01",
        current_mssv="2100001234",
        current_name="Nguyễn Văn A"
    )
    print("Kết quả phát hiện:")
    print(json.dumps(fraud_res_1, ensure_ascii=False, indent=2))
    assert fraud_res_1["is_fraud"] == True
    assert fraud_res_1["student_b"]["mssv"] == "2100009876"
    print("➔ KẾT QUẢ: Phát hiện chính xác Sinh viên B là nguồn gốc file!")

    # 5. TEST TÌNH HUỐNG GIAN LẬN 2: SV A mở file mình nhưng Copy-Paste nội dung từ file của B
    print("\n--- TEST GIAN LẬN 2: SV A COPY - PASTE NỘI DUNG TỪ BÀI B ---")
    # Giả lập: file của A bị dán nội dung từ file của B chứa thủy ấn số của B
    doc_a = Document(file_a)
    stolen_paragraph = doc_a.add_paragraph()
    stolen_text_with_b_watermark = encode_hidden_mssv_to_text(
        "Thực hiện Kế hoạch tổ chức kỳ thi kết thúc học phần của trường;",
        "2100009876" # MSSV của Sinh viên B
    )
    stolen_paragraph.add_run(stolen_text_with_b_watermark)
    file_a_tampered = "storage/templates/01_2100001234_tampered.docx"
    doc_a.save(file_a_tampered)

    fraud_res_2 = check_anti_cheat(
        docx_path=file_a_tampered,
        current_sbd="01",
        current_mssv="2100001234",
        current_name="Nguyễn Văn A"
    )
    print("Kết quả phát hiện thủy ấn vô hình:")
    print(json.dumps(fraud_res_2, ensure_ascii=False, indent=2))
    assert fraud_res_2["is_fraud"] == True
    assert fraud_res_2["student_b"]["mssv"] == "2100009876"
    print("➔ KẾT QUẢ: Bắt quả tang chính xác thủy ấn số vô hình của Sinh viên B!")

    # 6. TEST BỘ THẨM ĐỊNH THỂ THỨC (RULE ENGINE Ô 1 - Ô 9) TRÊN FILE THÔ
    print("\n--- TEST CHẤM ĐIỂM THỂ THỨC TRÊN FILE THÔ (CHƯA CHỈNH SỬA) ---")
    validator = Decree30Validator(file_a, "thong_bao_nhiem_vu", tmpl)
    report = validator.validate_document()
    print(f"Tổng tiêu chí đạt: {report['passed_count']} / {report['total_count']} ({report['percentage']}%)")
    print(f"Có được phép nộp bài không: {report['can_submit']} (Phải là False vì chưa sửa)")
    assert report["can_submit"] == False

    print("\n" + "=" * 60)
    print("TẤT CẢ CÁC BÀI TEST ĐÃ VƯỢT QUA 100% THÀNH CÔNG!")
    print("=" * 60)

if __name__ == "__main__":
    test_full_pipeline()
