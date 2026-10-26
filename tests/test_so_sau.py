import os
import docx
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from engine.validator import Decree30Validator

def test_so_sau():
    doc = Document()
    
    # 1. Căn lề chuẩn: Trên 2cm, Dưới 2cm, Trái 3cm, Phải 1.5cm
    section = doc.sections[0]
    section.top_margin = Inches(2.0 / 2.54)
    section.bottom_margin = Inches(2.0 / 2.54)
    section.left_margin = Inches(3.0 / 2.54)
    section.right_margin = Inches(1.5 / 2.54)
    section.page_width = Inches(21.0 / 2.54)
    section.page_height = Inches(29.7 / 2.54)
    
    # Header Table 2x2
    table_header = doc.add_table(rows=2, cols=2)
    table_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Cell 0, 0: Cơ quan ban hành (Số sau: 13pt)
    c00 = table_header.cell(0, 0)
    p_cq_top = c00.paragraphs[0]
    p_cq_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_top = p_cq_top.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO")
    r_top.font.name = "Times New Roman"
    r_top.font.size = Pt(13)
    
    p_cq = c00.add_paragraph()
    p_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cq = p_cq.add_run("TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH")
    r_cq.font.name = "Times New Roman"
    r_cq.font.size = Pt(13)
    r_cq.bold = True
    
    # Cell 0, 1: Quốc hiệu & Tiêu ngữ (Số sau: Quốc hiệu 13pt, Tiêu ngữ 14pt)
    c01 = table_header.cell(0, 1)
    p_qh = c01.paragraphs[0]
    p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_qh = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    r_qh.font.name = "Times New Roman"
    r_qh.font.size = Pt(13)
    r_qh.bold = True
    
    p_tn = c01.add_paragraph()
    p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tn = p_tn.add_run("Độc lập – Tự do – Hạnh phúc")
    r_tn.font.name = "Times New Roman"
    r_tn.font.size = Pt(14)
    r_tn.bold = True
    
    # Cell 1, 0: Số ký hiệu (Số sau: 13pt)
    c10 = table_header.cell(1, 0)
    p_so = c10.paragraphs[0]
    p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_so = p_so.add_run("Số: 125/TB-NTTU")
    r_so.font.name = "Times New Roman"
    r_so.font.size = Pt(13)
    
    # Cell 1, 1: Địa danh, ngày tháng (Số sau: 14pt)
    c11 = table_header.cell(1, 1)
    p_dd = c11.paragraphs[0]
    p_dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dd = p_dd.add_run("Thành phố Hồ Chí Minh, ngày 20 tháng 10 năm 2026")
    r_dd.font.name = "Times New Roman"
    r_dd.font.size = Pt(14)
    r_dd.italic = True
    
    # Dòng trống
    doc.add_paragraph()
    
    # Ô 5a & 5b (Tên loại & Trích yếu ngoài bảng: Số sau: 14pt)
    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_type = p_type.add_run("THÔNG BÁO")
    r_type.font.name = "Times New Roman"
    r_type.font.size = Pt(14)
    r_type.bold = True
    
    p_ty = doc.add_paragraph()
    p_ty.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ty = p_ty.add_run("Về việc tổ chức hội thảo khoa học cấp trường năm 2026")
    r_ty.font.name = "Times New Roman"
    r_ty.font.size = Pt(14)
    r_ty.bold = True
    
    # Ô 6: Nội dung (Số sau: 14pt, first line indent, justify, space after 6pt)
    body_texts = [
        "Căn cứ Kế hoạch hoạt động khoa học và công nghệ năm học 2025 - 2026 của Trường Đại học Nguyễn Tất Thành đã được phê duyệt, Phòng Khoa học và Công nghệ xin thông báo đến các đơn vị về việc tổ chức Hội thảo khoa học cấp trường.",
        "Mục tiêu của hội thảo nhằm tạo diễn đàn trao đổi học thuật, công bố các công trình nghiên cứu khoa học mới nhất của cán bộ, giảng viên và sinh viên, đồng thời thúc đẩy văn hóa nghiên cứu sáng tạo trong toàn trường.",
        "Đề nghị các Khoa, Viện, Trung tâm và các đơn vị trực thuộc thông báo rộng rãi đến toàn thể giảng viên, nhà nghiên cứu và sinh viên tích cực viết bài và tham dự hội thảo đúng thời hạn quy định."
    ]
    for b in body_texts:
        p_b = doc.add_paragraph()
        p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_b.paragraph_format.first_line_indent = Inches(1.27 / 2.54)
        p_b.paragraph_format.space_after = Pt(6)
        p_b.paragraph_format.line_spacing = 1.2
        r_b = p_b.add_run(b)
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(14)
        
    doc.add_paragraph()
    
    # Footer Table 1x2
    table_footer = doc.add_table(rows=1, cols=2)
    table_footer.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Cell 0, 0: Nơi nhận
    cell_nn = table_footer.cell(0, 0)
    p_nn_t = cell_nn.paragraphs[0]
    r_nn_t = p_nn_t.add_run("Nơi nhận:")
    r_nn_t.font.name = "Times New Roman"
    r_nn_t.font.size = Pt(12)
    r_nn_t.bold = True
    r_nn_t.italic = True
    
    items = ["- Ban Giám hiệu (để b/c);", "- Các Khoa, Viện;", "- Lưu: VT, KHCN."]
    for it in items:
        p_it = cell_nn.add_paragraph()
        p_it.paragraph_format.space_after = Pt(0)
        p_it.paragraph_format.line_spacing = 1.0
        r_it = p_it.add_run(it)
        r_it.font.name = "Times New Roman"
        r_it.font.size = Pt(11)
        
    # Cell 0, 1: Ký tên (Số sau: 14pt)
    cell_sig = table_footer.cell(0, 1)
    p_title = cell_sig.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("HIỆU TRƯỞNG")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.bold = True
    
    for _ in range(3):
        p_space = cell_sig.add_paragraph()
        p_space.paragraph_format.space_after = Pt(0)
        p_space.paragraph_format.line_spacing = 1.0
        
    p_name = cell_sig.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run("Trần Hoàng Nam")
    r_name.font.name = "Times New Roman"
    r_name.font.size = Pt(14)
    r_name.bold = True
    
    os.makedirs("storage/templates", exist_ok=True)
    out_path = "storage/templates/test_so_sau.docx"
    doc.save(out_path)
    print("✔ Đã tạo file chuẩn Số sau:", out_path)
    
    v = Decree30Validator(out_path, "thong_bao_nhiem_vu")
    res = v.validate_document()
    
    print("\n--- KẾT QUẢ KIỂM THỬ FILE CHUẨN SỐ SAU ---")
    print("Font Profile:", res.get("font_profile"))
    print(f"Tiêu chí đạt: {res['passed_count']} / {res['total_count']} ({res['percentage']}%)")
    print(f"Đủ điều kiện nộp bài: {res['can_submit']}")
    for item in res["checklist"]:
        status = "✔ ĐẠT" if item["passed"] else "✖ LỖI"
        print(f"  [{status}] {item['name']}")
        for err in item.get("errors", []):
            print(f"       - Lỗi: {err['issue']}")
            
    assert res["font_profile"] == "SO_SAU", "Font profile phải là SO_SAU"
    assert res["can_submit"] == True, "Phải đủ điều kiện nộp bài 100%"
    print("\n🎉 THÀNH CÔNG: Quy ước 'Số sau' hoạt động hoàn hảo 100%!")

if __name__ == "__main__":
    test_so_sau()
