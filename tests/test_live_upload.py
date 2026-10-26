import requests
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. Sinh viên tạo và tải file mẫu chính thức từ hệ thống
gen_res = requests.post("http://127.0.0.1:5000/api/template/generate", json={
    "class_code": "KN01",
    "sbd": "01",
    "mssv": "2100001234",
    "name": "Nguyễn Văn A",
    "doc_type": "thong_bao_nhiem_vu"
}).json()

filename = gen_res["filename"]
print("1. Generated student template:", filename)

# 2. Sinh viên mở file này trên Word và căn chỉnh chuẩn 100% Nghị định 30
# (Đảm bảo giữ nguyên metadata và thủy ấn bảo mật của file)
file_path = f"storage/templates/{filename}"
doc = Document(file_path)

# Xóa các đoạn text thô cũ trong body để đưa vào 2 bảng và nội dung chuẩn
for p in list(doc.paragraphs):
    p._element.getparent().remove(p._element)

# Căn lề A4 chuẩn NĐ 30: Trên 2cm, Dưới 2cm, Trái 3cm, Phải 1.5cm
for s in doc.sections:
    s.top_margin = Inches(2.0 / 2.54)
    s.bottom_margin = Inches(2.0 / 2.54)
    s.left_margin = Inches(3.0 / 2.54)
    s.right_margin = Inches(1.5 / 2.54)

# Bảng 1: Header 2x2
t_head = doc.add_table(rows=2, cols=2)
# Cell 0, 0: Cơ quan
p_cq = t_head.cell(0, 0).paragraphs[0]
p_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cq.paragraph_format.line_spacing = 1.0
p_cq.paragraph_format.space_after = Pt(0)
r1 = p_cq.add_run("TRƯỜNG ĐẠI HỌC NGUYỄN TẤT THÀNH\n")
r1.font.name = "Times New Roman"
r1.font.size = Pt(12)
r2 = p_cq.add_run("PHÒNG ĐÀO TẠO")
r2.font.name = "Times New Roman"
r2.font.size = Pt(12)
r2.bold = True

# Cell 0, 1: Quốc hiệu & Tiêu ngữ
p_qh = t_head.cell(0, 1).paragraphs[0]
p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_qh.paragraph_format.line_spacing = 1.0
p_qh.paragraph_format.space_after = Pt(0)
rqh = p_qh.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
rqh.font.name = "Times New Roman"
rqh.font.size = Pt(12)
rqh.bold = True

p_tn = t_head.cell(0, 1).add_paragraph()
p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_tn.paragraph_format.line_spacing = 1.0
p_tn.paragraph_format.space_before = Pt(0)
rtn = p_tn.add_run("Độc lập - Tự do - Hạnh phúc")
rtn.font.name = "Times New Roman"
rtn.font.size = Pt(13)
rtn.bold = True

# Cell 1, 0: Số ký hiệu
p_so = t_head.cell(1, 0).paragraphs[0]
p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
rso = p_so.add_run("Số: 45/TB-ĐT")
rso.font.name = "Times New Roman"
rso.font.size = Pt(13)

# Cell 1, 1: Địa danh ngày tháng
p_dd = t_head.cell(1, 1).paragraphs[0]
p_dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
rdd = p_dd.add_run("Thành phố Hồ Chí Minh, ngày 02 tháng 02 năm 2026")
rdd.font.name = "Times New Roman"
rdd.font.size = Pt(13)
rdd.italic = True

# Ô 5a: Tên loại & Trích yếu (ngoài bảng)
p_type = doc.add_paragraph()
p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_type.paragraph_format.space_before = Pt(12)
p_type.paragraph_format.space_after = Pt(4)
rtype = p_type.add_run("THÔNG BÁO")
rtype.font.name = "Times New Roman"
rtype.font.size = Pt(13)
rtype.bold = True

p_try = doc.add_paragraph()
p_try.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_try.paragraph_format.space_after = Pt(12)
rtry = p_try.add_run("Về việc phân công nhiệm vụ cán bộ coi thi kết thúc học phần")
rtry.font.name = "Times New Roman"
rtry.font.size = Pt(13)
rtry.bold = True

# Ô 6: Thân bài (Justify, thụt 1.27cm, giãn đoạn 6pt)
p_body = doc.add_paragraph()
p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_body.paragraph_format.first_line_indent = Inches(1.27 / 2.54)
p_body.paragraph_format.space_after = Pt(6)
p_body.paragraph_format.line_spacing = 1.15
rbody = p_body.add_run("Thực hiện Kế hoạch tổ chức kỳ thi kết thúc học phần học kỳ 1 năm học 2025 - 2026 đã được Ban Giám hiệu phê duyệt, Phòng Đào tạo thông báo đến toàn thể giảng viên, cán bộ quản lý trong Nhà trường về việc thực hiện nhiệm vụ coi thi học kỳ nghiêm túc theo đúng quy chế.")
rbody.font.name = "Times New Roman"
rbody.font.size = Pt(13)

# Bảng 2: Footer 1x2
t_foot = doc.add_table(rows=1, cols=2)
# Cell 0, 0: Nơi nhận
p_nn = t_foot.cell(0, 0).paragraphs[0]
rnn = p_nn.add_run("Nơi nhận:")
rnn.font.name = "Times New Roman"
rnn.font.size = Pt(12)
rnn.bold = True
rnn.italic = True

p_i1 = t_foot.cell(0, 0).add_paragraph()
p_i1.paragraph_format.space_after = Pt(0)
ri1 = p_i1.add_run("- Ban Giám hiệu (để báo cáo);")
ri1.font.name = "Times New Roman"
ri1.font.size = Pt(11)

p_i2 = t_foot.cell(0, 0).add_paragraph()
p_i2.paragraph_format.space_after = Pt(0)
ri2 = p_i2.add_run("- Các Khoa, Viện đào tạo;")
ri2.font.name = "Times New Roman"
ri2.font.size = Pt(11)

p_i3 = t_foot.cell(0, 0).add_paragraph()
p_i3.paragraph_format.space_after = Pt(0)
ri3 = p_i3.add_run("- Lưu: VT, ĐT.")
ri3.font.name = "Times New Roman"
ri3.font.size = Pt(11)

# Cell 0, 1: Ký tên
p_cv = t_foot.cell(0, 1).paragraphs[0]
p_cv.alignment = WD_ALIGN_PARAGRAPH.CENTER
rcv1 = p_cv.add_run("TL. HIỆU TRƯỞNG\nTRƯỞNG PHÒNG ĐÀO TẠO")
rcv1.font.name = "Times New Roman"
rcv1.font.size = Pt(13)
rcv1.bold = True

for _ in range(3):
    t_foot.cell(0, 1).add_paragraph()

p_name = t_foot.cell(0, 1).add_paragraph()
p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
rname = p_name.add_run("Nguyễn Văn An")
rname.font.name = "Times New Roman"
rname.font.size = Pt(13)
rname.bold = True

doc.save(file_path)
print("2. Student edited and formatted document successfully!")

# 3. Tải bài lên để thẩm định thể thức qua API /api/validate
with open(file_path, "rb") as f:
    files = {"file": (filename, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {
        "class_code": "KN01",
        "sbd": "01",
        "mssv": "2100001234",
        "name": "Nguyễn Văn A",
        "doc_type": "thong_bao_nhiem_vu"
    }
    r = requests.post("http://127.0.0.1:5000/api/validate", files=files, data=data)
    res = r.json()
    print("3. Validation Status:", res.get("status"))
    print(f"   Score: {res.get('percentage')}% ({res.get('passed_count')}/{res.get('total_count')})")
    print("   Can submit:", res.get("can_submit"))
    assert res.get("can_submit") == True, "Bài làm phải đạt chuẩn 100%!"

# 4. Xác nhận nộp bài qua API /api/submit
sub_res = requests.post("http://127.0.0.1:5000/api/submit", json={
    "class_code": "KN01",
    "sbd": "01",
    "mssv": "2100001234",
    "name": "Nguyễn Văn A",
    "doc_type": "thong_bao_nhiem_vu",
    "temp_file": res["temp_file"]
}).json()

print("4. Submission response:", sub_res.get("success"), "| Filename:", sub_res.get("submission", {}).get("filename"))
assert sub_res.get("success") == True, "Nộp bài phải thành công!"

print("\n🎉 HOÀN TẤT TOÀN DIỆN VÒNG ĐỜI KHẢO THÍ TỪ SINH VIÊN ĐẾN GIẢNG VIÊN VỚI ĐIỂM SỐ 100%!")
