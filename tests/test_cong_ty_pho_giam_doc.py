import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from engine.validator import Decree30Validator

def create_cong_van_cong_ty(docx_path, kt_text="KT.GIÁM ĐỐC"):
    doc = Document()
    
    # 1. Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # 2.0 cm
        section.bottom_margin = Inches(0.79)  # 2.0 cm
        section.left_margin = Inches(1.18)    # 3.0 cm
        section.right_margin = Inches(0.79)   # 2.0 cm

    # 2. Header Table (2x2)
    t_head = doc.add_table(rows=2, cols=2)
    t_head.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Cell 0, 0: Cơ quan ban hành là CÔNG TY
    p_cq = t_head.cell(0, 0).paragraphs[0]
    p_cq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cq = p_cq.add_run('CÔNG TY CỔ PHẦN CÔNG NGHỆ Á CHÂU')
    r_cq.font.name = 'Times New Roman'
    r_cq.font.size = Pt(12)
    r_cq.bold = True
    
    # Cell 0, 1: Quốc hiệu & Tiêu ngữ
    p_qh = t_head.cell(0, 1).paragraphs[0]
    p_qh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rqh = p_qh.add_run('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM')
    rqh.font.name = 'Times New Roman'
    rqh.font.size = Pt(12)
    rqh.bold = True
    
    p_tn = t_head.cell(0, 1).add_paragraph()
    p_tn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rtn = p_tn.add_run('Độc lập - Tự do - Hạnh phúc')
    rtn.font.name = 'Times New Roman'
    rtn.font.size = Pt(13)
    rtn.bold = True
    
    # Cell 1, 0: Số ký hiệu & Ô 5b Trích yếu
    p_so = t_head.cell(1, 0).paragraphs[0]
    p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_so.paragraph_format.space_after = Pt(6)
    rso = p_so.add_run('Số: 28/AC-VP')
    rso.font.name = 'Times New Roman'
    rso.font.size = Pt(13)
    
    p_5b = t_head.cell(1, 0).add_paragraph()
    p_5b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_5b.paragraph_format.space_before = Pt(6)
    r5b = p_5b.add_run('V/v đề nghị hợp tác phát triển phần mềm')
    r5b.font.name = 'Times New Roman'
    r5b.font.size = Pt(12)
    
    # Cell 1, 1: Địa danh ngày tháng
    p_dd = t_head.cell(1, 1).paragraphs[0]
    p_dd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rdd = p_dd.add_run('Thành phố Hồ Chí Minh, ngày 23 tháng 10 năm 2026')
    rdd.font.name = 'Times New Roman'
    rdd.font.size = Pt(13)
    rdd.italic = True
    
    # 3. Kính gửi (Ô 9a)
    p_kg = doc.add_paragraph()
    rkg = p_kg.add_run('Kính gửi: Sở Khoa học và Công nghệ Thành phố Hồ Chí Minh.')
    rkg.font.name = 'Times New Roman'
    rkg.font.size = Pt(13)
    
    # 4. Thân bài (Ô 6)
    p_body = doc.add_paragraph()
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body.paragraph_format.first_line_indent = Inches(0.5)
    p_body.paragraph_format.line_spacing = 1.15
    p_body.paragraph_format.space_after = Pt(6)
    rb = p_body.add_run('Công ty Cổ phần Công nghệ Á Châu xin gửi lời chào trân trọng đến Quý cơ quan. Nhằm đẩy mạnh ứng dụng chuyển đổi số trong giáo dục, Công ty trân trọng đề nghị Quý cơ quan xem xét hỗ trợ triển khai dự án.')
    rb.font.name = 'Times New Roman'
    rb.font.size = Pt(13)
    
    # 5. Footer Table (1x2)
    t_foot = doc.add_table(rows=1, cols=2)
    t_foot.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Cell 0, 0: Nơi nhận Ô 9b
    p_nn = t_foot.cell(0, 0).paragraphs[0]
    rnn = p_nn.add_run('Nơi nhận:')
    rnn.font.name = 'Times New Roman'
    rnn.font.size = Pt(12)
    rnn.bold = True
    rnn.italic = True
    
    p_i1 = t_foot.cell(0, 0).add_paragraph()
    ri1 = p_i1.add_run('- Như kính gửi;')
    ri1.font.name = 'Times New Roman'
    ri1.font.size = Pt(11)
    
    p_i2 = t_foot.cell(0, 0).add_paragraph()
    ri2 = p_i2.add_run('- Ban Giám đốc (để báo cáo);')
    ri2.font.name = 'Times New Roman'
    ri2.font.size = Pt(11)
    
    p_i3 = t_foot.cell(0, 0).add_paragraph()
    ri3 = p_i3.add_run('- Lưu: VT, VP.')
    ri3.font.name = 'Times New Roman'
    ri3.font.size = Pt(11)
    
    # Cell 0, 1: Ô 7a & 7b: Chức vụ KT. GIÁM ĐỐC / PHÓ GIÁM ĐỐC & Họ tên
    cell_sign = t_foot.cell(0, 1)
    p_cv1 = cell_sign.paragraphs[0]
    p_cv1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rcv1 = p_cv1.add_run(kt_text)
    rcv1.font.name = 'Times New Roman'
    rcv1.font.size = Pt(13)
    rcv1.bold = True
    
    p_cv2 = cell_sign.add_paragraph()
    p_cv2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rcv2 = p_cv2.add_run('PHÓ GIÁM ĐỐC')
    rcv2.font.name = 'Times New Roman'
    rcv2.font.size = Pt(13)
    rcv2.bold = True
    
    for _ in range(3):
        cell_sign.add_paragraph()
        
    p_name = cell_sign.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rname = p_name.add_run('Trần Thị Mai Phương')
    rname.font.name = 'Times New Roman'
    rname.font.size = Pt(13)
    rname.bold = True
    
    doc.save(docx_path)

def test_company_pho_giam_doc():
    print("=== KIỂM THỬ BÀI TẬP CÔNG TY: PHÓ GIÁM ĐỐC KÝ THAY (KT.GIÁM ĐỐC / PHÓ GIÁM ĐỐC) ===")
    
    # Test 1: KT.GIÁM ĐỐC (viết liền sau dấu chấm như trong ảnh chụp sinh viên)
    test_file_1 = "test_cong_ty_lien.docx"
    create_cong_van_cong_ty(test_file_1, kt_text="KT.GIÁM ĐỐC")
    
    validator1 = Decree30Validator(test_file_1, "cong_van_de_nghi", {})
    res1 = validator1.validate_document()
    
    sign_box1 = next((item for item in res1["checklist"] if item["box_id"] == 7), None)
    print(f"1. Kết quả kiểm tra Ô 7 & 8 với 'KT.GIÁM ĐỐC': Passed = {sign_box1['passed']}")
    if not sign_box1["passed"]:
        print("   Lỗi phát hiện:", sign_box1["errors"])
    assert sign_box1["passed"] == True, f"Ô 7 & 8 phải ĐẠT nhưng bị lỗi: {sign_box1['errors']}"
    assert res1["can_submit"] == True, f"Bài làm phải đạt 100% (Hiện tại: {res1['percentage']}%)"
    print(" [+] Test 1 ĐẠT 100% TIÊU CHUẨN!")
    
    # Test 2: KT. GIÁM ĐỐC (có khoảng trắng sau dấu chấm)
    test_file_2 = "test_cong_ty_cach.docx"
    create_cong_van_cong_ty(test_file_2, kt_text="KT. GIÁM ĐỐC")
    
    validator2 = Decree30Validator(test_file_2, "cong_van_de_nghi", {})
    res2 = validator2.validate_document()
    
    sign_box2 = next((item for item in res2["checklist"] if item["box_id"] == 7), None)
    print(f"2. Kết quả kiểm tra Ô 7 & 8 với 'KT. GIÁM ĐỐC': Passed = {sign_box2['passed']}")
    assert sign_box2["passed"] == True, f"Ô 7 & 8 phải ĐẠT nhưng bị lỗi: {sign_box2['errors']}"
    assert res2["can_submit"] == True, f"Bài làm phải đạt 100% (Hiện tại: {res2['percentage']}%)"
    print(" [+] Test 2 ĐẠT 100% TIÊU CHUẨN!")
    
    # Dọn dẹp file test
    for f in [test_file_1, test_file_2]:
        if os.path.exists(f):
            os.remove(f)
            
    print("\n🎉 XÁC NHẬN HOÀN TOÀN: BÀI TẬP CÔNG TY CÓ PHÓ GIÁM ĐỐC KÝ THAY (KT. GIÁM ĐỐC) ĐẠT 100% CHUẨN NGHỊ ĐỊNH 30!")

if __name__ == "__main__":
    test_company_pho_giam_doc()
